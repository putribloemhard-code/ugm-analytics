"""Generator konten Dampak & SDG (.docx) -- python-docx + kaleido.

Kaleido dipakai KHUSUS di sini untuk export chart Plotly ke gambar statis
(PNG) yang disisipkan ke dokumen Word -- docs/FRAMEWORK.md bilang "BUKAN
matplotlib/kaleido" untuk *laporan HTML statis* (laporan_static.py, yang
memang tidak butuh gambar raster karena JS Plotly-nya inline). Word TIDAK
bisa menampilkan chart interaktif, jadi kasusnya beda -- satu-satunya cara
menaruh chart di dokumen Word adalah sebagai gambar. Dikonfirmasi user.

Dua cara pakai:
- generate(...) -> dokumen BERDIRI SENDIRI (cover + ringkasan + pilar + SDG + catatan).
- append_lampiran(docx_bytes, ...) -> tempel bagian "Lampiran: Data Dampak &
  SDG" (ringkasan + pilar + SDG + catatan, TANPA cover baru) ke BELAKANG
  dokumen .docx yang sudah ada (dipakai page_akreditasi.py utk menyatukan
  dengan dokumen LED/LKPS dari subproyek akreditasi/).
"""

import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from kepmen_sdg import WARNA_PILAR, sdg_label  # noqa: E402
from unit_kerja import UNIT_KERJA  # noqa: E402


def _label_fakultas(fakultas_pilih: str) -> str:
    return fakultas_pilih if fakultas_pilih == "Seluruh Universitas" else UNIT_KERJA[fakultas_pilih]["nama"]


def _png_terlalu_gelap(png_bytes: bytes) -> bool:
    """Heuristik deteksi capture rusak: kaleido 1.x (browser-driven, via
    choreographer) kadang-kadang menangkap frame SEBELUM warna marker/colorbar
    selesai di-paint (race condition, teramati di mesin ini terutama saat
    sistem sedang berat -- lihat catatan _fig_to_stream) -- hasilnya semua
    bar chart tampak nyaris hitam alih-alih warna aslinya. Chart normal
    didominasi area putih/terang (sumbu, background, grid), jadi proporsi
    piksel nyaris-hitam yang tinggi adalah sinyal kuat capture rusak, bukan
    warna sah (bahkan skala warna paling gelap sekalipun -- mis. Viridis --
    tidak pernah mendominasi >25% kanvas)."""
    from io import BytesIO as _BIO

    from PIL import Image

    im = Image.open(_BIO(png_bytes)).convert("RGB").resize((60, 40))
    pixels = list(im.getdata())
    gelap = sum(1 for r, g, b in pixels if r < 40 and g < 40 and b < 60)
    # Dikalibrasi dari sampel nyata: capture rusak (semua bar nyaris hitam)
    # -> ~12% piksel gelap; chart normal (termasuk yang punya bar gelap sah,
    # mis. navy di skala Plasma) -> ~0%. Ambang 3% kasih margin aman di antaranya.
    return (gelap / len(pixels)) > 0.03


def _fig_to_stream(fig, width=900, height=460, scale=2, retries=2) -> BytesIO:
    """Export fig ke PNG; retry kalau hasilnya kedeteksi rusak (lihat
    _png_terlalu_gelap) -- render ulang biasanya langsung normal lagi karena
    ini race condition transien, bukan kesalahan konfigurasi chart."""
    png_bytes = None
    for _attempt in range(retries + 1):
        png_bytes = fig.to_image(format="png", width=width, height=height, scale=scale)
        if not _png_terlalu_gelap(png_bytes):
            break
    return BytesIO(png_bytes)


def _add_tabel_df(doc: Document, df: pd.DataFrame) -> None:
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr[i].text = str(col)
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            cells[i].text = str(row[col])


def _write_body(
    doc: Document,
    *,
    b: pd.DataFrame,
    bk_f: pd.DataFrame,
    bs_f: pd.DataFrame,
    rp_f: pd.DataFrame,
    tahun_awal: str,
    tahun_akhir: str,
    pilar_pilih: list[str],
    ringkasan: dict,
    catatan_metodologi: str,
) -> None:
    """Isi bersama: ringkasan eksekutif + pilar + SDG + catatan metodologi.
    Dipanggil oleh generate() (dokumen baru) dan append_lampiran() (ditempel
    ke dokumen lain) -- TIDAK menulis cover, itu tanggung jawab caller."""
    import plotly.express as px
    import plotly.io as pio

    # `import streamlit` diam-diam mengganti pio.templates.default dari
    # "plotly" jadi "streamlit" (dipakai Streamlit supaya chart LIVE di
    # dashboard ikut tema-nya) -- tapi kaleido (static export ke PNG utk
    # Word) gagal resolve continuous colorscale template itu, hasilnya
    # SEMUA bar chart warna-kontinu (mis. chart SDG di bawah, color="sdg")
    # jadi nyaris hitam alih-alih gradien warna aslinya. Root cause
    # dikonfirmasi lewat perbandingan langsung: sama persis kecuali baris
    # ini, satu render hitam total, satu gradien benar. HARUS di-reset
    # SEBELUM px.bar() dipanggil (di figure CREATION time) -- reset setelah
    # lewat fig.update_layout(template=...) TERBUKTI tidak cukup, warna
    # sudah kadung "dibakar" ke color axis saat px.bar() jalan. Disimpan &
    # dikembalikan di finally supaya chart LIVE di dashboard (di luar fungsi
    # ini, proses Streamlit yang sama) tidak ikut kehilangan tema Streamlit.
    _template_asal = pio.templates.default
    pio.templates.default = "plotly"
    try:
        _write_body_charts(doc, b=b, bk_f=bk_f, bs_f=bs_f, rp_f=rp_f, tahun_awal=tahun_awal,
                            tahun_akhir=tahun_akhir, pilar_pilih=pilar_pilih, ringkasan=ringkasan,
                            catatan_metodologi=catatan_metodologi, px=px)
    finally:
        pio.templates.default = _template_asal


def _write_body_charts(
    doc: Document, *, b, bk_f, bs_f, rp_f, tahun_awal, tahun_akhir, pilar_pilih, ringkasan,
    catatan_metodologi: str, px,
) -> None:
    doc.add_heading("Ringkasan Eksekutif", level=1)
    rc = doc.add_table(rows=1, cols=4)
    rc.style = "Light Grid Accent 1"
    labels = ["Total berita dampak", "Dampak pertumbuhan tertinggi", f"Sorotan {tahun_akhir}",
              ringkasan["topik_top_kind_label"]]
    values = [f"{ringkasan['total_berita']:,}", ringkasan["pilar_top"],
              f"{ringkasan['berita_tahun_ini']:,} berita", ringkasan["topik_top_short"]]
    for i, l in enumerate(labels):
        rc.rows[0].cells[i].text = l
        for p in rc.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
    row2 = rc.add_row().cells
    for i, v in enumerate(values):
        row2[i].text = v
    doc.add_paragraph()
    doc.add_paragraph(ringkasan["narasi"])
    doc.add_page_break()

    doc.add_heading("Dampak per Pilar (Lingkungan / Ekonomi / Sosial)", level=1)
    if len(rp_f):
        fig_rp = px.bar(
            rp_f, x="dampak", y="jumlah_berita", color="dampak",
            title="Berita unik per dampak (semua tema Kepmen)",
            labels={"dampak": "Dampak", "jumlah_berita": "Jumlah berita"},
            color_discrete_map=WARNA_PILAR,
        )
        fig_rp.update_layout(showlegend=False, height=420, width=800)
        doc.add_picture(_fig_to_stream(fig_rp), width=Inches(6))
        tabel_pilar = rp_f[["dampak", "jumlah_berita"]].rename(
            columns={"dampak": "Dampak", "jumlah_berita": "Jumlah Berita"}
        )
        _add_tabel_df(doc, tabel_pilar)
    else:
        doc.add_paragraph("Tidak ada data dampak untuk filter ini.")
    doc.add_page_break()

    doc.add_heading("Keterkaitan dengan SDGs", level=1)
    if len(bs_f):
        dist_s = (
            bs_f.groupby("sdg")["url"].nunique().reset_index(name="jumlah").sort_values("jumlah", ascending=False)
        )
        dist_s["label"] = dist_s["sdg"].map(lambda s: f"SDG {s}")
        dist_s["nama"] = dist_s["sdg"].apply(sdg_label)
        fig_s = px.bar(
            dist_s.sort_values("jumlah"), x="label", y="jumlah", color="sdg",
            title="Jumlah berita per SDG (klaster resmi)",
            labels={"label": "SDG", "jumlah": "Jumlah berita"},
        )
        fig_s.update_layout(showlegend=False, height=420, width=800,
                            xaxis=dict(tickangle=-45, tickfont=dict(size=10)))
        doc.add_picture(_fig_to_stream(fig_s), width=Inches(6))
        tabel_sdg = dist_s[["label", "nama", "jumlah"]].rename(
            columns={"label": "SDG", "nama": "Nama", "jumlah": "Jumlah Berita"}
        )
        _add_tabel_df(doc, tabel_sdg)
    else:
        doc.add_paragraph("Tidak ada data SDG untuk filter ini.")
    doc.add_page_break()

    doc.add_heading("Catatan Metodologi & Batasan", level=1)
    for para in catatan_metodologi.strip().split("\n\n"):
        doc.add_paragraph(para.strip())


def generate(*, b, bk_f, bs_f, rp_f, tahun_awal, tahun_akhir, pilar_pilih, ringkasan,
             catatan_metodologi: str, fakultas_pilih: str = "Seluruh Universitas") -> bytes:
    """Dokumen berdiri sendiri: cover + ringkasan + pilar + SDG + catatan."""
    doc = Document()
    title = doc.add_heading("Laporan Akreditasi — Analisis Dampak UGM", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Sumber: berita ugm.ac.id — Kepmendikti Saintek 361/M/KEP/2025 & 17 SDGs").italic = True
    p_tgl = doc.add_paragraph()
    p_tgl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tgl.add_run(f"Digenerate: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph()
    p_filter = doc.add_paragraph()
    p_filter.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_filter.add_run(
        f"Filter — Rentang tahun: {tahun_awal}–{tahun_akhir}  |  "
        f"Dampak: {', '.join(pilar_pilih) if pilar_pilih else 'Semua'}  |  "
        f"Fakultas: {_label_fakultas(fakultas_pilih)}"
    )
    doc.add_page_break()

    _write_body(doc, b=b, bk_f=bk_f, bs_f=bs_f, rp_f=rp_f, tahun_awal=tahun_awal,
                tahun_akhir=tahun_akhir, pilar_pilih=pilar_pilih, ringkasan=ringkasan,
                catatan_metodologi=catatan_metodologi)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def append_lampiran(docx_bytes: bytes, *, b, bk_f, bs_f, rp_f, tahun_awal, tahun_akhir,
                     pilar_pilih, ringkasan, catatan_metodologi: str,
                     fakultas_pilih: str = "Seluruh Universitas") -> bytes:
    """Tempel bagian "Lampiran: Data Dampak & SDG" ke BELAKANG dokumen .docx
    yang sudah ada (mis. hasil akreditasi/scripts/generate_template.py)."""
    doc = Document(BytesIO(docx_bytes))
    doc.add_page_break()
    doc.add_heading("Lampiran: Data Dampak & SDG (Berita-Dampak)", level=1)
    filter_p = doc.add_paragraph()
    filter_p.add_run(
        f"Filter — Rentang tahun: {tahun_awal}–{tahun_akhir}  |  "
        f"Dampak: {', '.join(pilar_pilih) if pilar_pilih else 'Semua'}  |  "
        f"Fakultas: {_label_fakultas(fakultas_pilih)}"
    ).italic = True

    _write_body(doc, b=b, bk_f=bk_f, bs_f=bs_f, rp_f=rp_f, tahun_awal=tahun_awal,
                tahun_akhir=tahun_akhir, pilar_pilih=pilar_pilih, ringkasan=ringkasan,
                catatan_metodologi=catatan_metodologi)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
