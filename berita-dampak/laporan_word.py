"""Generator laporan Word (.docx) ringkasan analisis Dampak/Dampak x SDGs/SDGs.

REUSE pola python-docx yang sudah terbukti di
akreditasi/scripts/generate_laporan_live.py (Document(), add_heading/
add_paragraph/add_table, simpan ke BytesIO) -- BUKAN library atau pola baru.

SENGAJA tidak memakai kaleido/matplotlib untuk merender chart Plotly jadi
gambar: proyek ini sudah punya pengalaman buruk dengan dependency itu di
Windows (lihat catatan di percakapan). Isi tiap chart disajikan sbg TABEL data
(dataframe yang sama yang dipakai chart-nya), ditambah kalimat insight yang
SAMA dengan yang tampil di layar (fungsi penjelasan()/st.info() di halaman) --
dihitung dari dataframe ter-filter, bukan AI, jadi laporan selalu akurat
mengikuti filter yang aktif saat tombol diklik.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

# Navy sama dengan identitas visual dashboard (lihat common.py / background.png).
WARNA_JUDUL = RGBColor(0x00, 0x21, 0x4A)
MAKS_BARIS_TABEL = 40  # tabel lebih panjang dipotong + catatan "dan N baris lainnya"


@dataclass
class SeksiLaporan:
    """Satu chart/tabel di halaman -> satu section di laporan Word: judul,
    dataframe (tabel data chart itu, BUKAN gambar chart), dan insight (teks
    yang sama dengan yang sudah tampil di layar via penjelasan()/st.info())."""

    judul: str
    df: pd.DataFrame | None = None
    insight: str = ""
    catatan: str = ""
    maks_baris: int = MAKS_BARIS_TABEL


def _tulis_markdown_sederhana(paragraf, teks: str) -> None:
    """Insight di layar pakai **tebal** (markdown Streamlit) -- pecah jadi run
    bold/biasa di Word, supaya tampilannya konsisten dengan di layar, bukan
    tampil apa adanya sbg tanda bintang literal."""
    for i, bagian in enumerate(re.split(r"\*\*(.+?)\*\*", teks)):
        run = paragraf.add_run(bagian)
        if i % 2 == 1:
            run.bold = True


def _format_sel(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int):
        return f"{v:,}"
    return "" if v is None else str(v)


def _tambah_tabel(doc: Document, df: pd.DataFrame | None, maks_baris: int) -> None:
    if df is None or not len(df):
        doc.add_paragraph("Tidak ada data untuk filter ini.").runs[0].italic = True
        return
    kolom = list(df.columns)
    table = doc.add_table(rows=1, cols=len(kolom))
    table.style = "Light Grid Accent 1"
    for i, k in enumerate(kolom):
        cell = table.rows[0].cells[i]
        cell.text = str(k)
        cell.paragraphs[0].runs[0].font.bold = True
    for _, row in df.head(maks_baris).iterrows():
        cells = table.add_row().cells
        for i, k in enumerate(kolom):
            cells[i].text = _format_sel(row[k])
    if len(df) > maks_baris:
        p = doc.add_paragraph(f"... dan {len(df) - maks_baris:,} baris lainnya (total {len(df):,} baris).")
        p.runs[0].italic = True


def _tambah_cover(doc: Document, judul: str, subjudul: str, filter_lines: list[str]) -> None:
    title = doc.add_heading(judul, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = WARNA_JUDUL
    if subjudul:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.add_run(subjudul).italic = True
    p_tgl = doc.add_paragraph()
    p_tgl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tgl.add_run(f"Digenerate: {datetime.now().strftime('%Y-%m-%d %H:%M')} WIB "
                  "-- dari data live MySQL sesuai filter aktif saat itu.").italic = True
    doc.add_heading("Filter yang Diterapkan", level=1)
    if filter_lines:
        for line in filter_lines:
            doc.add_paragraph(line, style="List Bullet")
    else:
        doc.add_paragraph("Tidak ada filter khusus (data ditampilkan penuh).")


def build_laporan(
    *,
    judul: str,
    subjudul: str,
    filter_lines: list[str],
    ringkasan_eksekutif: str,
    seksi: list[SeksiLaporan],
) -> bytes:
    """Rangkai satu laporan .docx: cover + filter aktif + ringkasan eksekutif +
    N section (tiap section = 1 chart/tabel di halaman: judul, tabel data,
    insight). Dipanggil dari tombol "Unduh Laporan (Word)" di halaman
    Dampak/Dampak x SDGs/SDGs -- lihat page_dampak.py & page_sdgs.py."""
    doc = Document()
    _tambah_cover(doc, judul, subjudul, filter_lines)

    doc.add_heading("Ringkasan Eksekutif", level=1)
    p = doc.add_paragraph()
    _tulis_markdown_sederhana(p, ringkasan_eksekutif)

    for s in seksi:
        doc.add_heading(s.judul, level=1)
        if s.catatan:
            doc.add_paragraph(s.catatan).runs[0].italic = True
        _tambah_tabel(doc, s.df, s.maks_baris)
        if s.insight:
            ip = doc.add_paragraph()
            ip.add_run("Insight: ").bold = True
            _tulis_markdown_sederhana(ip, s.insight)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def tombol_unduh_laporan(
    *,
    key: str,
    judul: str,
    subjudul: str,
    filter_lines: list[str],
    ringkasan_eksekutif: str,
    seksi: list[SeksiLaporan],
    nama_file_bagian: list[str],
) -> None:
    """UI lengkap "Unduh Laporan (Word)": caption + tombol generate + tombol
    unduh. Pola generate->download SAMA seperti "Generate Dokumen" di
    akreditasi/dashboard_akreditasi.py (bytes dibuat di dalam if st.button,
    st.download_button langsung ditaruh di situ juga -- TIDAK butuh
    session_state, Streamlit sudah mengirim bytes-nya ke browser pada render
    yang sama). Satu fungsi dipakai bareng oleh page_dampak.py (2x -- sebelum
    & sesudah drill-down pilar) dan page_sdgs.py supaya UI-nya konsisten."""
    st.markdown("---")
    st.subheader("📄 Unduh Laporan")
    st.caption(
        "Dokumen Word (.docx) berisi filter aktif, ringkasan eksekutif, dan tabel data + "
        "insight dari tiap chart yang sedang tampil di halaman ini (bukan gambar chart -- "
        "datanya dalam bentuk tabel)."
    )
    if st.button("🔄 Buat Laporan (Word)", key=f"buat_{key}"):
        with st.spinner("Membuat laporan..."):
            docx_bytes = build_laporan(
                judul=judul, subjudul=subjudul, filter_lines=filter_lines,
                ringkasan_eksekutif=ringkasan_eksekutif, seksi=seksi,
            )
        st.download_button(
            "⬇️ Unduh Laporan (Word)",
            data=docx_bytes,
            file_name=nama_file(*nama_file_bagian),
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key=f"dl_{key}",
        )


def nama_file(*bagian: str) -> str:
    """"Laporan", "Dampak", "Lingkungan", "2024" -> Laporan_Dampak_Lingkungan_2024.docx
    (spasi & karakter non-alfanumerik jadi "_", supaya nama file selalu aman
    dipakai di Windows)."""
    potongan = []
    for b in bagian:
        b = re.sub(r"[^\w-]+", "_", str(b), flags=re.UNICODE).strip("_")
        if b:
            potongan.append(b)
    return "_".join(potongan) + ".docx"
