"""Uji status sumber resmi (data_source_map.json), data live, pratinjau ekstraksi, dan Word gabungan."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import text

from app.services import accreditation_sumber as sumber
from test_accreditation_workspace import EMAIL, _item, buat_laporan, engine, lkps, svc  # noqa: F401  (fixture)

LIVE_DDL = [
    """CREATE TABLE akreditasi_item_tersedia (item_id VARCHAR(64), baris_ke INTEGER, kolom VARCHAR(255), nilai TEXT,
       sumber_url VARCHAR(1000), fetched_at TIMESTAMP)""",
    """CREATE TABLE akreditasi_publikasi_dosen (link VARCHAR(500), dosen VARCHAR(255), sinta_id VARCHAR(32), platform VARCHAR(32),
       judul TEXT, tahun INTEGER, sumber VARCHAR(255), fetched_at TIMESTAMP)""",
    """CREATE TABLE akreditasi_berita_dcse (url VARCHAR(500), judul TEXT, tanggal VARCHAR(32), deskripsi TEXT, kategori VARCHAR(64),
       relevan_mei INTEGER, fetched_at TIMESTAMP)""",
]


@pytest.fixture()
def live(engine):  # noqa: F811
    now = datetime(2026, 9, 11, 13, 8)
    with engine.begin() as conn:
        for ddl in LIVE_DDL:
            conn.execute(text(ddl))
        conn.execute(text("INSERT INTO akreditasi_item_tersedia VALUES ('lkps_6', 1, 'Visi PT', 'Visi UGM', 'https://ugm.ac.id/visi', :t)"),
                     {"t": now})
        conn.execute(text("INSERT INTO akreditasi_item_tersedia VALUES ('lkps_6', 1, 'Kolom Tambahan', 'x', 'https://ugm.ac.id/visi', :t)"),
                     {"t": now})
        conn.execute(text("INSERT INTO akreditasi_publikasi_dosen VALUES ('https://scopus/1', 'Dosen A', '1', 'scopus', 'Judul A', 2025, 'Jurnal', :t)"),
                     {"t": now})
        conn.execute(text("INSERT INTO akreditasi_berita_dcse VALUES ('https://dcse/1', 'Berita A', '2026-01-02', '', 'x', 1, :t)"),
                     {"t": now})
    return engine


def test_status_mengikuti_peta_sumber_bukan_registry():
    # 61 item di peta: 5 tersedia, 41 akses data, 15 perlu disusun tim (lihat data_source_map.json).
    kategori = [sumber.kategori(i) for i in sumber.peta()]
    assert (kategori.count("tersedia"), kategori.count("akses_data"), kategori.count("penyusunan")) == (5, 41, 15)
    assert sumber.kategori("lkps_6") == "tersedia" and sumber.info_item("lkps_6")["sumber_asli"]


def test_tanpa_tabel_live_halaman_tetap_jalan(svc, lkps):  # noqa: F811
    w = svc.workspace(lkps)
    item = _item(w, "lkps_6")
    assert item["live"] is None and item["state"] == "kosong" and item["kategori"] == "tersedia"
    assert w["ringkasan"]["tersedia"] == {"total": 2, "lengkap": 0}


def test_data_live_dihitung_lengkap_dan_bisa_diganti_isian_tim(svc, live, lkps):  # noqa: F811
    item = _item(svc.workspace(lkps), "lkps_6")
    assert item["state"] == "live" and item["live"]["rows"][0]["Visi PT"] == "Visi UGM"
    assert item["live"]["kolom"][-1] == "Kolom Tambahan" and item["live"]["sumber"] == ["https://ugm.ac.id/visi"]
    assert svc.workspace(lkps)["ringkasan"]["tersedia"]["lengkap"] == 1
    pend = _item(svc.workspace(lkps), "lkps_3_c_2")["pendukung"]
    assert pend["total"] == 1 and pend["tautan"] == ["https://scopus/1"]
    svc.save_item(lkps, "lkps_6", [{item["kolom"][0]: "Isian tim"}], EMAIL)
    assert _item(svc.workspace(lkps), "lkps_6")["state"] == "terisi"


def test_data_live_hanya_untuk_prodi_sumbernya(svc, live):  # noqa: F811
    with live.begin() as conn:
        conn.execute(text("INSERT INTO akreditasi_prodi (id, fakultas_id, nama, jenjang, slug) VALUES (2, 1, 'Ilmu Komputer', 'Sarjana', 's1-ilkom')"))
    item = _item(svc.workspace(buat_laporan(live, "LKPS", prodi="s1-ilkom")), "lkps_6")
    assert item["live"] is None and item["pendukung"] is None


def test_word_gabungan_isian_tim_live_dan_placeholder(svc, live, lkps):  # noqa: F811
    kolom = _item(svc.workspace(lkps), "lkps_3_c_1")["kolom"]
    svc.save_item(lkps, "lkps_3_c_1", [{kolom[0]: "Kerja sama resmi"}], EMAIL)
    content, _ = svc.generate(lkps, EMAIL)
    doc = Document(BytesIO(content))
    teks = "\n".join(p.text for p in doc.paragraphs)
    sel = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Magister Elektronika dan Instrumentasi" in teks
    assert "Kerja sama resmi" in sel and "Sumber: isian tim penyusun." in teks
    assert "Visi UGM" in teks and "Sumber live: https://ugm.ac.id/visi" in teks
    assert "[DATA TIDAK TERSEDIA]" in teks and "Judul A" in sel  # placeholder akses data + data pendukung SINTA
    assert "tersedia live 1/2" in teks
