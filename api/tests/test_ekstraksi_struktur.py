"""Ekstraksi tanpa AI (akreditasi/scripts/ekstraksi_struktur.py) pada dokumen Word sintetis.

PDF LED asli prodi divalidasi terpisah oleh akreditasi/scripts/validasi_ekstraksi_struktur.py (butuh
berkas referensi di docs/sumber); uji ini menjaga jalur Word, kata kunci fallback, dan pemetaan kolom.
"""
from __future__ import annotations

import pytest
from docx import Document

from app.domain.source import load_accreditation_module

es = load_accreditation_module("ekstraksi_struktur.py")


@pytest.fixture()
def led_docx(tmp_path):
    d = Document()
    d.add_heading("IDENTITAS TIM PENYUSUN", 1)
    t = d.add_table(rows=3, cols=4)
    for r, isi in enumerate([["1", "Nama", ":", "Dr. Rina Lestari"], ["", "NIDN", ":", "0011"],
                             ["", "Jabatan", ":", "Ketua Tim"]]):
        for c, v in enumerate(isi):
            t.cell(r, c).text = v
    d.add_heading("KATA PENGANTAR", 1)
    d.add_paragraph("Paragraf pertama kata pengantar.")
    d.add_paragraph("Paragraf kedua, ditutup salam.")
    # Judul tidak persis ("... Laporan") -> lewat kata kunci fallback.
    d.add_heading("Ringkasan Eksekutif Laporan", 1)
    d.add_paragraph("Isi ringkasan.")
    d.add_heading("BAB I. PENDAHULUAN", 1)
    d.add_heading("A. Dasar Penyusunan", 2)
    d.add_paragraph("Dasar hukum penyusunan.")
    d.add_heading("C. Mekanisme Kerja Penyusunan Evaluasi Diri", 2)
    d.add_paragraph("Tahapan kerja tim.")
    d.add_heading("Kriteria C", 1)
    d.add_heading("C1. BUDAYA MUTU", 2)
    d.add_heading("C1.1 [PENETAPAN] Dokumen kebijakan budaya mutu", 3)
    d.add_paragraph("Kebijakan ditetapkan dalam Renstra.")
    d.add_paragraph("Tabel C1.1 Indikator sistem tata kelola UPPS terkait administrasi akademik")
    tb = d.add_table(rows=4, cols=7)
    head = ["No.", "Referensi", "Indikator", "IKU/IKT", "Sasaran/Target", "", ""]
    for c, v in enumerate(head):
        tb.cell(0, c).text = v
    tb.cell(0, 4).merge(tb.cell(0, 6))
    for c, v in enumerate(["", "", "", "", "TS-2", "TS-1", "TS"]):
        tb.cell(1, c).text = v
    for r, isi in enumerate([["1.", "Renstra hal. 12", "Sistem informasi terintegrasi", "IKT", "2", "2", "3"],
                             ["2.", "Renstra hal. 22", "Indeks kepuasan layanan", "IKU", "3,5", "3,6", "3,7"]], start=2):
        for c, v in enumerate(isi):
            tb.cell(r, c).text = v
    # Tahap tanpa nomor/kurung siku: hanya kata kunci "pelaksanaan" di dalam bab C1.
    p = d.add_paragraph()
    p.add_run("Pelaksanaan").bold = True
    d.add_paragraph("Program kerja dijalankan tiap semester.")
    d.add_heading("BAB III. PENUTUP", 1)
    d.add_paragraph("Penutup laporan.")
    path = tmp_path / "led-lama.docx"
    d.save(path)
    return str(path)


def test_word_led_lama_disalin_utuh(led_docx):
    rows, ring = es.ekstrak_dokumen(led_docx, "docx", "LED")
    assert ring["error"] is None and ring["metode"] == "tanpa_ai"
    nilai = {(r["item_id"], r["baris_ke"], r["nama_kolom"]): r["nilai"] for r in rows}
    assert nilai[("kata_pengantar", 1, "Kata Pengantar")] == "Paragraf pertama kata pengantar.\n\nParagraf kedua, ditutup salam."
    assert nilai[("ringkasan_eksekutif", 1, "Ringkasan Eksekutif")] == "Isi ringkasan."
    assert nilai[("bab1_pendahuluan", 1, "A. Dasar Penyusunan")] == "Dasar hukum penyusunan."
    assert nilai[("bab1_pendahuluan", 1, "B. Mekanisme Kerja Penyusunan Evaluasi Diri")] == "Tahapan kerja tim."
    assert nilai[("bab3_penutup", 1, "Bab III Penutup")] == "Penutup laporan."
    # Narasi PPEPP: teks bab tanpa isi tabel; tahap kedua ketemu lewat kata kunci.
    penetapan = nilai[("led_c1_budaya_mutu", 1, "Dokumen Penetapan / Kebijakan (link)")]
    assert "Kebijakan ditetapkan dalam Renstra." in penetapan and "Sistem informasi" not in penetapan
    assert nilai[("led_c1_budaya_mutu", 1, "Bukti Pelaksanaan")] == "Program kerja dijalankan tiap semester."
    # Tabel: header bertingkat -> kolom registry; satu record per baris bernomor.
    kolom = es.KEBUTUHAN_DATA["led_tabel_c1_1"]["kolom_dibutuhkan"]
    tabel = {k: v for (i, b, k), v in nilai.items() if i == "led_tabel_c1_1" and b == 2}
    assert tabel == dict(zip(kolom, ["Renstra hal. 22", "Indeks kepuasan layanan", "IKU", "3,5", "3,6", "3,7"]))
    assert nilai[("tim_penyusun_led", 1, "Nama")] == "Dr. Rina Lestari"
    assert nilai[("tim_penyusun_led", 1, "Jabatan dalam Tim")] == "Ketua Tim"


@pytest.mark.parametrize("header, harapan", [
    (["No.", "Referensi", "Indikator", "IKU/ IKT", "Sasaran/Target – TS-2", "Sasaran/Target – TS-1", "Sasaran/Target – TS"], 1.0),
    # Sel header tergabung (empat label jadi satu) -> kolom deskriptif tidak ketemu -> ditolak.
    (["IKU/ No Referensi Indikator IKT", "Target – TS-2", "Target – TS-1", "Target – TS"], 0.0),
    # Label ganda = header rusak.
    (["Indikator", "Target – TS-2", "Target – TS-2"], 0.0),
])
def test_skor_header_menolak_header_rusak(header, harapan):
    assert es._skor(header, es.KEBUTUHAN_DATA["led_tabel_c1_1"]["kolom_dibutuhkan"]) == harapan


def test_ts_polos_tidak_dipetakan_ke_semester():
    # Registry C2.2 meminta per semester (TS-2 – Ganjil/Genap); header dokumen hanya TS-2/TS-1/TS.
    kolom = es.KEBUTUHAN_DATA["led_tabel_c2_2"]["kolom_dibutuhkan"]
    assert es.petakan_kolom(["Parameter", "TS-2", "TS-1", "TS"], kolom) == {"Parameter": "Parameter"}


def _tabel_ts(ws, baris: int, judul: str | None, data: list[list]) -> int:
    """Tulis [caption] + header 2 tingkat (Sasaran/Target di atas TS-2/TS-1/TS) + data mulai `baris`."""
    if judul:
        ws.cell(baris, 1, judul)
        baris += 1
    for c, v in enumerate(["No.", "Referensi", "Indikator", "IKU/", "Sasaran/Target"], start=1):
        ws.cell(baris, c, v)
    ws.merge_cells(start_row=baris, start_column=5, end_row=baris, end_column=7)
    for c in (1, 2, 3):
        ws.merge_cells(start_row=baris, start_column=c, end_row=baris + 1, end_column=c)
    for c, v in enumerate(["IKT", "TS-2", "TS-1", "TS"], start=4):
        ws.cell(baris + 1, c, v)
    for i, r in enumerate(data, start=baris + 2):
        for c, v in enumerate(r, start=1):
            ws.cell(i, c, v)
    return baris + 2 + len(data) + 1   # + satu baris kosong pemisah


def test_excel_rekap_bebas_per_caption_dan_nama_sheet(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap LED"
    b = _tabel_ts(ws, 1, "Tabel C1.1 Indikator sistem tata kelola UPPS terkait administrasi akademik",
                  [[1, "Renstra hal. 12", "Sistem informasi", "IKT", 2, 2, 3.0]])
    # Tabel tanpa caption di sheet umum: kolomnya identik C1.2-C1.6, jadi TIDAK ditebak.
    _tabel_ts(ws, b, None, [[1, "Renstra", "Tidak boleh masuk mana pun", "IKU", 9, 9, 9]])
    ws2 = wb.create_sheet("C1.2")   # nama sheet = kode tabel
    _tabel_ts(ws2, 1, None, [[1, "Renstra hal. 30", "Akreditasi internasional", "IKU", 1, 2, 2],
                             [2, "Renstra hal. 31", "Prodi unggul", "IKU", 4, 4, 5]])
    path = tmp_path / "rekap.xlsx"
    wb.save(path)

    rows, ring = es.ekstrak_dokumen(str(path), "xlsx", "LED")
    assert ring["error"] is None
    per = {}
    for r in rows:
        per.setdefault(r["item_id"], {}).setdefault(r["baris_ke"], {})[r["nama_kolom"]] = r["nilai"]
    assert set(per) == {"led_tabel_c1_1", "led_tabel_c1_2"}
    assert per["led_tabel_c1_1"][1] == {"Referensi": "Renstra hal. 12", "Indikator": "Sistem informasi", "IKU/IKT": "IKT",
                                        "Sasaran/Target – TS-2": "2", "Sasaran/Target – TS-1": "2", "Sasaran/Target – TS": "3"}
    assert [r["Indikator"] for r in per["led_tabel_c1_2"].values()] == ["Akreditasi internasional", "Prodi unggul"]
    assert "Tidak boleh masuk mana pun" not in {r["nilai"] for r in rows}
