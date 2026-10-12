"""Draf AI narasi LED (services/accreditation_narasi_ai.py) dengan klien LLM tiruan dan SQLite sementara."""
from __future__ import annotations

import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from sqlalchemy import text

from app.services.accreditation_narasi_ai import DIISI_AI, _parse, buat_draf, item_draf_ai, pulihkan
from app.services.accreditation_workspace import NotFound, WorkspaceError
from tests.test_accreditation_workspace import EMAIL, TABEL, _item, buat_laporan, engine, led, lkps, svc  # noqa: F401

NARASI = "led_c3_penelitian"


class KlienTiruan:
    def __init__(self, jawaban=None, galat: Exception | None = None):
        self.panggilan: list[list[dict]] = []
        self.jawaban, self.galat = jawaban, galat
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

    def _create(self, model: str, messages: list[dict]):
        self.panggilan.append(messages)
        if self.galat:
            raise self.galat
        kolom = json.loads(messages[1]["content"].split("KOLOM YANG DIMINTA (JSON key): ")[1].split("\n")[0])
        isi = self.jawaban or {k: f"Draf {k}. [perlu dilengkapi: dokumen {k}]" for k in kolom}
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(isi)))])


def test_daftar_item_draf_ai_hanya_narasi_led_tanpa_identitas():
    ids = item_draf_ai()
    assert "identitas_pt_upps_ps" not in ids and NARASI in ids and TABEL not in ids
    assert len(ids) == 19


def test_buat_draf_tersimpan_bertanda_draf_ai_dan_memakai_data_lkps(svc, engine, led, lkps):
    k = _item(svc.workspace(lkps), TABEL)["kolom"]
    svc.save_item(lkps, TABEL, [{k[0]: "Penelitian bersama BRIN 2025"}], EMAIL)
    klien = KlienTiruan()
    hasil = buat_draf(engine, led, NARASI, client=klien)
    assert hasil["penanda"] == 5
    prompt = klien.panggilan[0][1]["content"]
    assert "Penelitian bersama BRIN 2025" in prompt and "Magister Elektronika dan Instrumentasi" in prompt

    item = _item(svc.workspace(led), NARASI)
    assert item["draf_ai"] is True and item["diisi_oleh"] == DIISI_AI and item["bisa_draf_ai"] is True
    assert item["rows"][0][item["kolom"][0]].startswith("Draf ")

    buat_draf(engine, led, NARASI, client=klien)  # draf AI boleh dibuat ulang
    svc.save_item(led, NARASI, item["rows"], EMAIL)  # disunting tim -> bukan draf lagi
    assert _item(svc.workspace(led), NARASI)["draf_ai"] is False


def test_draf_ai_tidak_menimpa_isian_tim_atau_item_final(svc, engine, led):
    k = _item(svc.workspace(led), NARASI)["kolom"]
    svc.save_item(led, NARASI, [{k[0]: "Tulisan tim"}], EMAIL)
    klien = KlienTiruan()
    with pytest.raises(WorkspaceError, match="sudah diisi tim"):
        buat_draf(engine, led, NARASI, client=klien)

    lain = "led_c4_pkm"
    svc.set_final(led, [lain], True, EMAIL)
    with pytest.raises(WorkspaceError, match="final"):
        buat_draf(engine, led, lain, client=klien)
    assert not klien.panggilan
    assert _item(svc.workspace(led), NARASI)["rows"][0][k[0]] == "Tulisan tim"


def test_timpa_mencadangkan_isian_tim_dan_bisa_dipulihkan(svc, engine, led):
    k = _item(svc.workspace(led), NARASI)["kolom"]
    svc.save_item(led, NARASI, [{k[0]: "Tulisan tim"}], EMAIL)
    klien = KlienTiruan()
    buat_draf(engine, led, NARASI, client=klien, timpa=True)
    buat_draf(engine, led, NARASI, client=klien)  # buat ulang draf: cadangan isian tim tetap
    item = _item(svc.workspace(led), NARASI)
    assert item["draf_ai"] and item["bisa_pulihkan"]

    assert pulihkan(engine, led, NARASI)["message"]
    item = _item(svc.workspace(led), NARASI)
    assert item["rows"][0][k[0]] == "Tulisan tim" and item["diisi_oleh"] == EMAIL
    assert not item["draf_ai"] and not item["bisa_pulihkan"]
    with pytest.raises(WorkspaceError, match="Tidak ada isian sebelumnya"):
        pulihkan(engine, led, NARASI)


def test_pulihkan_ditolak_bila_draf_sudah_disunting(svc, engine, led):
    k = _item(svc.workspace(led), NARASI)["kolom"]
    svc.save_item(led, NARASI, [{k[0]: "Tulisan tim"}], EMAIL)
    buat_draf(engine, led, NARASI, client=KlienTiruan(), timpa=True)
    svc.save_item(led, NARASI, [{k[0]: "Draf yang sudah disunting"}], EMAIL)
    assert not _item(svc.workspace(led), NARASI)["bisa_pulihkan"]
    with pytest.raises(WorkspaceError, match="sudah disunting"):
        pulihkan(engine, led, NARASI)
    assert _item(svc.workspace(led), NARASI)["rows"][0][k[0]] == "Draf yang sudah disunting"


def test_draf_ai_ditolak_untuk_lkps_identitas_dan_item_asing(engine, led, lkps):
    klien = KlienTiruan()
    with pytest.raises(WorkspaceError):
        buat_draf(engine, lkps, TABEL, client=klien)
    with pytest.raises(WorkspaceError):
        buat_draf(engine, led, "identitas_pt_upps_ps", client=klien)
    with pytest.raises(NotFound):
        buat_draf(engine, led, "item_tidak_ada", client=klien)
    assert not klien.panggilan


def test_galat_llm_tidak_menyimpan_apa_pun(svc, engine, led):
    with pytest.raises(WorkspaceError, match="tidak bisa dihubungi"):
        buat_draf(engine, led, NARASI, client=KlienTiruan(galat=TimeoutError("habis")))
    with engine.connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM akreditasi_data_manual")).scalar() == 0


def test_kolom_yang_dikosongkan_ai_diberi_penanda(svc, engine, led):
    kolom = _item(svc.workspace(led), NARASI)["kolom"]
    buat_draf(engine, led, NARASI, client=KlienTiruan({kolom[0]: "Isi pertama."}))
    rows = _item(svc.workspace(led), NARASI)["rows"][0]
    assert rows[kolom[0]] == "Isi pertama." and rows[kolom[1]] == f"[perlu dilengkapi: {kolom[1]}]"


def test_parse_jawaban_berpagar_dan_kolom_hilang():
    assert _parse('```json\n{"A": " isi ", "X": "asing"}\n```', ["A", "B"]) == {"A": "isi", "B": ""}
    assert _parse('Berikut draf: {"A": "isi"} selesai', ["A"]) == {"A": "isi"}
    with pytest.raises(WorkspaceError):
        _parse("bukan json", ["A"])


def test_word_menandai_draf_ai_wajib_diperiksa(svc, engine, led):
    buat_draf(engine, led, NARASI, client=KlienTiruan())
    content, _ = svc.generate(led, EMAIL)
    teks = "\n".join(p.text for p in Document(BytesIO(content)).paragraphs)
    assert "Draf AI yang Wajib Diperiksa" in teks
    assert "Sumber: DRAF AI dari data laporan" in teks

    svc.set_final(led, [NARASI], True, EMAIL)  # sudah direview -> tidak lagi masuk daftar periksa
    teks = "\n".join(p.text for p in Document(BytesIO(svc.generate(led, EMAIL)[0])).paragraphs)
    assert "Draf AI yang Wajib Diperiksa" not in teks
