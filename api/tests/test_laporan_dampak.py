"""Uji laporan dampak otomatis (services/laporan_dampak.py): kerangka laporan resmi, angka, pratinjau = Word."""
from __future__ import annotations

import base64
import io
import json
from datetime import date

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.domain.models import FilterParams
from app.services.laporan_dampak import build_laporan, id_angka, render_docx
from app.services.sources import ringkasan_sumber
from app.services.story import build_story
from test_story import make_frames, make_matkul


def frames():
    fr = make_frames()
    fr.matkul = make_matkul()
    return fr


def laporan(mode: str, **filter_kw):
    fr = frames()
    story = build_story(fr, FilterParams(**filter_kw), mode, matkul=fr.matkul)
    return build_laporan(story, ringkasan_sumber(fr), hari_ini=date(2026, 9, 24))


def teks(lap) -> str:
    return "\n".join(b.get("text", "") + " ".join(b.get("items", [])) for b in lap["blocks"])


def test_kerangka_dampak_mengikuti_daftar_isi_laporan_resmi():
    lap = laporan("impact")
    judul = [e["text"] for e in lap["toc"]]
    urutan = ["RINGKASAN EKSEKUTIF", "LEMBAR IDENTIFIKASI", "BAB I - PENDAHULUAN", "1.1 Latar Belakang", "1.2 Tujuan",
              "1.3 Ruang Lingkup", "1.4 Unit Analisis", "BAB II - DAMPAK SOSIAL", "2.1 Pendidikan Inklusif",
              "BAB III - DAMPAK EKONOMI", "3.3 Ekosistem Kewirausahaan", "BAB IV - DAMPAK LINGKUNGAN", "4.1 Energi",
              "BAB V - KESIMPULAN", "REFERENSI", "LAMPIRAN"]
    posisi = [judul.index(j) for j in urutan]
    assert posisi == sorted(posisi)
    assert lap["title"] == "LAPORAN DAMPAK SOSIAL, EKONOMI, DAN LINGKUNGAN"
    assert lap["generated"] == "24 September 2026"


def test_angka_laporan_berasal_dari_data_dan_jujur_soal_batasnya():
    isi = teks(laporan("impact"))
    # u4 & u5 = 2 berita Pendidikan Inklusif pada data uji
    assert "memuat 2 berita yang berkaitan dengan tema Pendidikan Inklusif" in isi
    assert "bukan nilai capaian indikator" in isi
    # tema tanpa berita tetap punya sub-bab dan dinyatakan kosong, bukan diberi angka
    assert "belum ditemukan berita UGM yang memuat kata kunci tema Transportasi" in isi
    identitas = next(b for b in laporan("impact")["blocks"] if b["type"] == "table")
    assert ["Koordinator dan pengesahan", "Diisi oleh unit penyusun sebelum laporan disahkan"] in identitas["rows"]


def test_gambar_dan_tabel_bernomor_berurutan():
    lap = laporan("impact-sdgs")
    gambar = [b["caption"] for b in lap["blocks"] if b["type"] == "figure"]
    tabel = [b["caption"] for b in lap["blocks"] if b["type"] == "table"]
    assert gambar and all(c.startswith(f"Gambar {i}. ") for i, c in enumerate(gambar, start=1))
    assert tabel and all(c.startswith(f"Tabel {i}. ") for i, c in enumerate(tabel, start=1))
    assert lap["figures"] == gambar and lap["tables"] == tabel
    # mode Dampak × SDGs menambah gambar sebaran SDG per tema
    assert any("per SDG" in c for c in gambar)
    png = base64.b64decode(next(b["image"] for b in lap["blocks"] if b["type"] == "figure"))
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_filter_pilar_hanya_memuat_bab_terpilih_dengan_nomor_resmi():
    lap = laporan("impact", pillars=("Lingkungan",))
    judul = [e["text"] for e in lap["toc"]]
    assert "BAB IV - DAMPAK LINGKUNGAN" in judul and "4.1 Energi" in judul and "BAB V - KESIMPULAN" in judul
    assert not any("SOSIAL" in j for j in judul)
    assert "Dampak: Lingkungan" in lap["filters"]


def test_laporan_sdgs_berisi_profil_per_sdg():
    lap = laporan("sdgs")
    judul = [e["text"] for e in lap["toc"]]
    assert "BAB II - SEBARAN PEMBERITAAN PER SDG" in judul
    assert "3.4 SDG 4: Pendidikan Berkualitas" in judul
    assert lap["title"].startswith("LAPORAN KONTRIBUSI UGM")
    json.dumps(lap)


def test_docx_memuat_isi_yang_sama_dengan_pratinjau():
    lap = laporan("impact")
    doc = Document(io.BytesIO(render_docx(lap)))
    isi_doc = "\n".join(p.text for p in doc.paragraphs)
    for b in lap["blocks"]:
        if b["type"] in ("heading", "paragraph"):
            assert b["text"] in isi_doc
    assert len(doc.inline_shapes) == len(lap["figures"])
    assert "DAFTAR ISI" in isi_doc and "DAFTAR GAMBAR" in isi_doc


def test_id_angka_format_indonesia():
    assert id_angka("naik 7,405 berita dari 12,534.") == "naik 7.405 berita dari 12.534."
    assert id_angka("56% lebih tinggi, SDG 1, 4") == "56% lebih tinggi, SDG 1, 4"


@pytest.mark.parametrize("mode", ["impact", "sdgs"])
def test_endpoint_pratinjau_dan_unduh(monkeypatch, mode):
    from app.api.v1 import analytics
    from app.main import app
    from app.services import story

    class Fake:
        engine = object()

    fr = frames()
    monkeypatch.setattr(story.StoryService, "frames", lambda self: fr)
    analytics._LAPORAN_CACHE.clear()
    sebelumnya = dict(app.dependency_overrides)
    app.dependency_overrides[analytics.service] = lambda: Fake()
    try:
        client = TestClient(app)
        r = client.post("/api/v1/analytics/reports/preview", json={"mode": mode})
        assert r.status_code == 200 and r.json()["blocks"]
        d = client.post("/api/v1/analytics/reports", json={"mode": mode})
        assert d.status_code == 200
        assert d.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert Document(io.BytesIO(d.content)).paragraphs
        assert client.post("/api/v1/analytics/reports", json={"mode": "lain"}).status_code == 422
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(sebelumnya)


def test_suntingan_paragraf_hanya_mengubah_narasi_dan_tidak_menyentuh_cache():
    from app.services.laporan_dampak import terapkan_suntingan
    lap = laporan("impact")
    i_par = next(i for i, b in enumerate(lap["blocks"]) if b["type"] == "paragraph")
    i_hapus = next(i for i, b in enumerate(lap["blocks"]) if b["type"] == "paragraph" and i > i_par)
    i_judul = next(i for i, b in enumerate(lap["blocks"]) if b["type"] == "heading")
    asli = lap["blocks"][i_par]["text"]
    baru = terapkan_suntingan(lap, {str(i_par): "Narasi hasil suntingan tim.", str(i_hapus): "  "})
    assert baru["blocks"][i_par]["text"] == "Narasi hasil suntingan tim."
    assert len(baru["blocks"]) == len(lap["blocks"]) - 1          # paragraf yang dikosongkan dihapus
    assert lap["blocks"][i_par]["text"] == asli                    # laporan asli (cache) utuh
    for salah in ({str(i_judul): "x"}, {"999999": "x"}, {"abc": "x"}, {str(i_par): "x" * 5001}):
        with pytest.raises(ValueError):
            terapkan_suntingan(lap, salah)
    doc = Document(io.BytesIO(render_docx(baru)))
    assert "Narasi hasil suntingan tim." in "\n".join(p.text for p in doc.paragraphs)
