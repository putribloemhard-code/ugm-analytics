"""Template Excel isian tabel akreditasi: unduh format, baca upload, dan migrasi kolom lama."""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from app.services import accreditation_tabel as tabel
from app.services import migrasi_kolom_tabel as migrasi

REG = tabel._registry()


def _isi(item_id: str, nilai: dict[int, list]) -> bytes:
    """Template item_id dengan baris data {indeks_baris: [nilai per kolom isian]}."""
    data, _ = tabel.buat_template(item_id)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    dalam, _ = tabel._sel_header(REG.FORMAT_TABEL[item_id])
    geser = 1 if REG.FORMAT_TABEL[item_id]["nomor"] else 0
    for i, baris in nilai.items():
        for j, v in enumerate(baris):
            ws.cell(2 + dalam + i, 1 + geser + j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_semua_isian_tabel_punya_format_dan_kolomnya_sama():
    tabel_ids = {k for k, v in REG.KEBUTUHAN_DATA.items() if v["tipe"] == "tabel"}
    assert set(REG.FORMAT_TABEL) == tabel_ids | {"identitas_pt_upps_ps"}  # + Identitas (tabel vertikal)
    assert len(tabel_ids) == 80
    for item_id, fmt in REG.FORMAT_TABEL.items():
        assert REG.KEBUTUHAN_DATA[item_id]["kolom_dibutuhkan"] == [REG.kunci_kolom(k) for k in fmt["kolom"]]


@pytest.mark.parametrize("item_id", sorted(k for k, v in REG.FORMAT_TABEL.items() if not tabel.vertikal(v)))
def test_template_dibaca_kembali_ke_kolom_yang_sama(item_id):
    kolom = REG.KEBUTUHAN_DATA[item_id]["kolom_dibutuhkan"]
    hasil = tabel.baca_upload(item_id, _isi(item_id, {0: [f"v{j}" for j in range(len(kolom))], 2: [12.0]}))
    assert hasil["kolom"] == kolom
    assert hasil["rows"][0] == {k: f"v{j}" for j, k in enumerate(kolom)}
    assert hasil["rows"][1][kolom[0]] == "12" and len(hasil["rows"]) == 2  # baris kosong dilewati


def test_file_tabel_lain_ditolak_walau_penanda_dihapus():
    isi = _isi("lkps_3_c_3", {0: ["HKI A"]})
    with pytest.raises(tabel.TabelError, match="bukan"):
        tabel.baca_upload("lkps_4_c_3", isi)  # kolom HKI penelitian & PkM sama persis
    wb = load_workbook(io.BytesIO(isi))
    wb.properties.keywords = None
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(tabel.TabelError, match="Judul"):
        tabel.baca_upload("lkps_4_c_3", buf.getvalue())


def test_header_diubah_ditolak():
    wb = load_workbook(io.BytesIO(_isi("led_tabel_c1_1", {0: ["Ref"]})))
    wb.active.cell(2, 2, "Kolom Lain")
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(tabel.TabelError, match="Header"):
        tabel.baca_upload("led_tabel_c1_1", buf.getvalue())


def test_migrasi_ganti_nama_arsip_dan_putar_visi():
    engine = create_engine("sqlite://")
    with engine.begin() as conn:
        conn.execute(text("""CREATE TABLE akreditasi_data_manual (id INTEGER PRIMARY KEY, prodi_id TEXT, item_id TEXT,
            baris_ke INTEGER, kolom TEXT, tahun INTEGER, nilai TEXT, link_bukti TEXT, diisi_oleh TEXT,
            updated_at TEXT, laporan_id INTEGER)"""))
        sel = [("lkps_d1", 1, "Nama MK", "Sensor"), ("lkps_d1", 1, "Alasan Kekhasan", "Inti"),
               ("lkps_1_a_5", 1, "D4/D3", "2"), ("lkps_3_c_3", 1, "Tahun Perolehan (TS-2/TS-1/TS)", "TS-1"),
               ("lkps_6", 1, "Level (PT/UPPS/PS)", "PT"), ("lkps_6", 1, "Teks Visi", "Visi UGM"),
               ("lkps_6", 1, "Teks Misi", "Misi UGM"), ("lkps_6", 2, "Level (PT/UPPS/PS)", "PS"),
               ("lkps_6", 2, "Teks Visi", "Visi Prodi")]
        for item_id, baris, kolom, nilai in sel:
            conn.execute(text("INSERT INTO akreditasi_data_manual (prodi_id, item_id, baris_ke, kolom, nilai, laporan_id) "
                              "VALUES ('mei', :i, :b, :k, :n, 7)"), {"i": item_id, "b": baris, "k": kolom, "n": nilai})

    laporan = migrasi.migrasi(engine)  # mode laporan: tidak mengubah apa pun
    assert laporan["ganti_manual"] == 2 and laporan["visi_manual"] == 2
    migrasi.migrasi(engine, terapkan=True)
    assert migrasi.migrasi(engine)["ganti_manual"] == 0  # idempoten

    with engine.connect() as conn:
        sel = {(r.item_id, r.baris_ke, r.kolom): r.nilai for r in conn.execute(text("SELECT * FROM akreditasi_data_manual"))}
        arsip = {r.kolom for r in conn.execute(text(f"SELECT kolom FROM {migrasi.ARSIP}"))}
    assert sel[("lkps_d1", 1, "Nama Mata Kuliah")] == "Sensor"
    assert sel[("lkps_d1", 1, "Keterangan Kekhasan")] == "Inti"
    assert sel[("lkps_3_c_3", 1, "Tahun Perolehan (beri tanda √) – TS-1")] == "√"
    assert sel[("lkps_6", 1, "Visi PT")] == "Visi UGM" and sel[("lkps_6", 1, "Visi Keilmuan PS")] == "Visi Prodi"
    assert {"D4/D3", "Teks Misi", "Level (PT/UPPS/PS)", "Tahun Perolehan (TS-2/TS-1/TS)"} <= arsip
    assert not any(k[2] in arsip for k in sel)


def test_identitas_gabungan_dipecah_per_butir():
    hasil = migrasi.pecah_identitas({
        "Nama & Alamat PT/UPPS/PS": "Perguruan Tinggi: UGM; Unit Pengelola Program Studi: DIKE; "
                                    "Program Studi: Magister EI; Alamat: Sekip Utara, Sleman",
        "Kontak": "Telepon: 0274 546194; E-mail: mei@ugm.ac.id; Website: https://x/mei/",
        "No. & Tanggal SK Pendirian PT": "PP No 23 Tahun 1949, tanggal 16 Desember 1949",
        "No. & Tanggal SK Pembukaan PS": "SK Rektor 211/2023; 16 Februari 2023",
        "Pejabat Penandatangan": "Rektor UGM",
    })
    assert hasil == {
        "Perguruan Tinggi": "UGM", "Unit Pengelola Program Studi": "DIKE", "Nama Program Studi": "Magister EI",
        "Alamat": "Sekip Utara, Sleman", "Nomor Telepon": "0274 546194",
        "E-mail dan Website": "mei@ugm.ac.id dan https://x/mei/",
        "Nomor SK Pendirian PT": "PP No 23 Tahun 1949", "Tanggal SK Pendirian PT": "16 Desember 1949",
        "Nomor SK Pembukaan PS": "SK Rektor 211/2023", "Tanggal SK Pembukaan PS": "16 Februari 2023",
        "Pejabat Penandatangan SK Pembukaan PS": "Rektor UGM",
    }
    assert set(hasil) <= set(REG.KEBUTUHAN_DATA["identitas_pt_upps_ps"]["kolom_dibutuhkan"])


def test_identitas_template_vertikal_satu_record():
    data, nama = tabel.buat_template("identitas_pt_upps_ps")
    assert nama.startswith("Format LED A. Spesifikasi Program")
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    kolom = REG.KEBUTUHAN_DATA["identitas_pt_upps_ps"]["kolom_dibutuhkan"]
    assert ws.cell(1, 1).value == "A. Spesifikasi Program" and [ws.cell(3 + i, 1).value for i in range(16)] == kolom
    ws.cell(3, 2, "Universitas Gadjah Mada")
    ws.cell(16, 2, 2023)
    buf = io.BytesIO()
    wb.save(buf)
    hasil = tabel.baca_upload("identitas_pt_upps_ps", buf.getvalue())
    assert len(hasil["rows"]) == 1
    assert hasil["rows"][0]["Perguruan Tinggi"] == "Universitas Gadjah Mada"
    assert hasil["rows"][0]["Tahun Pertama Kali Menerima Mahasiswa"] == "2023"

    ws.cell(4, 1, "Fakultas")  # butir diubah -> ditolak
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(tabel.TabelError, match="A4"):
        tabel.baca_upload("identitas_pt_upps_ps", buf.getvalue())
    with pytest.raises(tabel.TabelError):
        tabel.baca_upload("lkps_1_a_1", data)  # template identitas ke tabel lain
