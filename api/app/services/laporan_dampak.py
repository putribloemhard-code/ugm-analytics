"""Laporan dampak otomatis mengikuti kerangka LAPORAN DAMPAK SOSIAL, EKONOMI, DAN LINGKUNGAN UGM 2025.

Satu model dokumen (daftar blok: judul, paragraf, daftar, gambar, tabel) dipakai dua kali:
dikirim sebagai JSON untuk pratinjau di web, lalu dirender ke .docx saat diunduh, sehingga isi
pratinjau sama dengan dokumen Word. Semua angka berasal dari payload `build_story` (filter aktif)
dan ringkasan sumber; narasi disusun dari angka itu, tanpa angka karangan. Data dasarnya adalah
pemberitaan dan kurikulum, jadi setiap sub-bab menyatakan bahwa angkanya jejak publikasi, bukan
nilai capaian indikator resmi.
"""
from __future__ import annotations

import base64
import re
from datetime import date
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
# API objek matplotlib (Figure) dipakai, bukan pyplot: pyplot menyimpan state global dan tidak
# aman dipanggil paralel dari threadpool FastAPI.
from matplotlib.artist import setp
from matplotlib.figure import Figure

NAVY = "#01416b"
PILAR_WARNA = {"Sosial": "#d95926", "Ekonomi": "#3987e5", "Lingkungan": "#199e70"}
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September",
         "Oktober", "November", "Desember"]
MODE_JUDUL = {"impact": "Dampak", "impact-sdgs": "Dampak × SDGs", "sdgs": "SDGs"}
PENGANTAR_PILAR = {
    "Sosial": ("Dampak sosial mengukur kontribusi perguruan tinggi dalam meningkatkan akses, pemerataan, "
               "pemberdayaan, dan kualitas hidup masyarakat, meliputi pendidikan inklusif, penelitian dan "
               "inovasi, pengabdian dan pengembangan masyarakat, serta kontribusi dalam mendukung kebijakan publik."),
    "Ekonomi": ("Dampak ekonomi mengukur kontribusi perguruan tinggi terhadap pertumbuhan ekonomi, hilirisasi "
                "hasil riset, kewirausahaan, dan penguatan ekonomi lokal, termasuk aktivitas mahasiswa dan "
                "pengunjung, pemanfaatan hasil riset, spin-off/start-up, serta belanja institusi kepada UMKM lokal."),
    "Lingkungan": ("Dampak lingkungan mengukur kontribusi perguruan tinggi dalam mendukung pembangunan "
                   "berkelanjutan dan pelestarian lingkungan, meliputi pengelolaan energi, konsumsi yang "
                   "bertanggung jawab, transportasi berkelanjutan, keanekaragaman hayati, serta pendidikan dan "
                   "penelitian yang mendukung keberlanjutan lingkungan."),
}
HURUF = "abcdefghijklmnopqrstuvwxyz"


def ang(value: Any) -> str:
    """Angka format Indonesia (titik pemisah ribuan)."""
    try:
        return f"{int(value):,}".replace(",", ".")
    except (TypeError, ValueError):
        return str(value)


def pct(bagian: float, total: float) -> str:
    return f"{100 * bagian / total:.1f}".replace(".", ",") + "%" if total else "0%"


def tanggal_id(d: date) -> str:
    return f"{d.day} {BULAN[d.month - 1]} {d.year}"


def _daftar_kata(items: list[str]) -> str:
    items = [i for i in items if i]
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " dan " + items[-1]


def id_angka(text: str | None) -> str:
    """Kalimat insight dashboard memakai 7,405; laporan memakai format Indonesia 7.405."""
    return re.sub(r"\d{1,3}(?:,\d{3})+", lambda m: m.group(0).replace(",", "."), text or "")


def _kalimat(text: str | None) -> str:
    text = id_angka(text).strip()
    if text and text[-1] not in ".!?":
        text += "."
    return text


# ------------------------------------------------------------------ gambar grafik
def _gaya(ax: Any) -> None:
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.spines["left"].set_color("#9aa7b4")
    ax.spines["bottom"].set_color("#9aa7b4")
    ax.tick_params(colors="#33414e", labelsize=8)
    ax.grid(axis="both", color="#e3e8ee", linewidth=0.6)
    ax.set_axisbelow(True)


def _png(fig: Any) -> str:
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=130, bbox_inches="tight", facecolor="white")
    return base64.b64encode(buf.getvalue()).decode("ascii")


def gambar_batang(labels: list[str], values: list[float], horizontal: bool = False,
                  warna: list[str] | str = NAVY, satuan: str = "berita") -> str:
    tinggi = max(2.4, 0.32 * len(labels) + 0.8) if horizontal else 3.0
    fig = Figure(figsize=(6.4, tinggi))
    ax = fig.subplots()
    if horizontal:
        pos = list(range(len(labels)))[::-1]
        bars = ax.barh(pos, values, color=warna, height=0.62)
        ax.set_yticks(pos, labels)
        ax.grid(axis="y", visible=False)
        ax.set_xlabel(f"Jumlah {satuan}", fontsize=8, color="#33414e")
        puncak = max(values) if values else 0
        for bar, v in zip(bars, values):
            ax.text(bar.get_width() + puncak * 0.01, bar.get_y() + bar.get_height() / 2, ang(v),
                    va="center", fontsize=7.5, color="#33414e")
        ax.set_xlim(0, puncak * 1.12 if puncak else 1)
    else:
        bars = ax.bar(labels, values, color=warna, width=0.66)
        ax.grid(axis="x", visible=False)
        ax.set_ylabel(f"Jumlah {satuan}", fontsize=8, color="#33414e")
        if len(labels) > 8:
            setp(ax.get_xticklabels(), rotation=45, ha="right")
        puncak = max(values) if values else 0
        if len(labels) <= 17:
            for bar, v in zip(bars, values):
                ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + puncak * 0.01, ang(v),
                        ha="center", va="bottom", fontsize=7, color="#33414e")
        ax.set_ylim(0, puncak * 1.12 if puncak else 1)
    _gaya(ax)
    return _png(fig)


def gambar_garis(xs: list[str], ys: list[float]) -> str:
    fig = Figure(figsize=(6.4, 2.7))
    ax = fig.subplots()
    ax.plot(xs, ys, color=NAVY, linewidth=2, marker="o", markersize=3.5)
    if ys:
        i = max(range(len(ys)), key=lambda k: ys[k])
        ax.annotate(ang(ys[i]), (xs[i], ys[i]), textcoords="offset points", xytext=(0, 6), ha="center",
                    fontsize=7.5, color="#33414e")
        ax.set_ylim(0, max(ys) * 1.18 or 1)
    if len(xs) > 10:
        setp(ax.get_xticklabels(), rotation=45, ha="right")
    ax.set_ylabel("Jumlah berita", fontsize=8, color="#33414e")
    _gaya(ax)
    return _png(fig)


def gambar_heatmap(rows: list[str], cols: list[str], values: list[list[float]]) -> str:
    fig = Figure(figsize=(6.6, max(2.2, 0.3 * len(rows) + 1.0)))
    ax = fig.subplots()
    im = ax.imshow(values, aspect="auto", cmap="Blues")
    ax.set_yticks(range(len(rows)), rows)
    step = 1 if len(cols) <= 12 else 2
    ax.set_xticks(range(0, len(cols), step), cols[::step])
    setp(ax.get_xticklabels(), rotation=45, ha="right")
    ax.tick_params(labelsize=7.5, colors="#33414e", length=0)
    for side in ax.spines.values():
        side.set_visible(False)
    cbar = fig.colorbar(im, ax=ax, fraction=0.03, pad=0.02)
    cbar.ax.tick_params(labelsize=7)
    cbar.set_label("Jumlah berita", fontsize=7.5)
    return _png(fig)


def _chart(charts: list[dict[str, Any]], suffix: str) -> dict[str, Any] | None:
    return next((c for c in charts if str(c.get("id", "")).endswith(suffix)), None)


# ------------------------------------------------------------------ penyusun dokumen
class _Laporan:
    def __init__(self) -> None:
        self.blocks: list[dict[str, Any]] = []
        self.n_gambar = 0
        self.n_tabel = 0
        self.daftar_isi: list[dict[str, Any]] = []
        self.daftar_gambar: list[str] = []
        self.daftar_tabel: list[str] = []

    def bab(self, text: str, masuk_daftar: bool = True) -> None:
        self.blocks.append({"type": "heading", "level": 1, "text": text, "break_before": True})
        if masuk_daftar:
            self.daftar_isi.append({"level": 1, "text": text})

    def sub(self, text: str, level: int = 2, masuk_daftar: bool = True) -> None:
        self.blocks.append({"type": "heading", "level": level, "text": text})
        if masuk_daftar and level == 2:
            self.daftar_isi.append({"level": 2, "text": text})

    def p(self, text: str) -> None:
        if text and text.strip():
            self.blocks.append({"type": "paragraph", "text": text.strip()})

    def daftar(self, items: list[str], bernomor: bool = False) -> None:
        items = [i for i in items if i]
        if items:
            self.blocks.append({"type": "list", "ordered": bernomor, "items": items})

    def gambar(self, image: str, caption: str, sumber: str | None = None) -> None:
        self.n_gambar += 1
        judul = f"Gambar {self.n_gambar}. {caption}"
        self.daftar_gambar.append(judul)
        self.blocks.append({"type": "figure", "image": image, "caption": judul, "source": sumber})

    def tabel(self, caption: str, columns: list[str], rows: list[list[Any]], links: dict[int, int] | None = None,
              sumber: str | None = None, kunci_nilai: bool = False) -> None:
        self.n_tabel += 1
        judul = f"Tabel {self.n_tabel}. {caption}"
        self.daftar_tabel.append(judul)
        self.blocks.append({"type": "table", "caption": judul, "columns": columns,
                            "rows": [["" if v is None else str(v) for v in r] for r in rows],
                            "links": {str(k): v for k, v in (links or {}).items()}, "source": sumber,
                            "key_value": kunci_nilai})


def _berita_sumber(sumber: dict[str, Any] | None) -> dict[str, Any]:
    return (sumber or {}).get("berita") or {}


def _latar_belakang(lap: _Laporan, mode: str) -> None:
    lap.p("Perguruan tinggi memiliki peran strategis dalam mendukung pembangunan berkelanjutan melalui pendidikan, "
          "penelitian, pengabdian kepada masyarakat, serta pengembangan inovasi yang memberikan manfaat nyata bagi "
          "masyarakat. Kontribusi tersebut tidak hanya tercermin dari capaian akademik, tetapi juga dari dampak sosial, "
          "ekonomi, dan lingkungan yang dihasilkan.")
    lap.p("Keputusan Menteri Pendidikan Tinggi, Sains, dan Teknologi Nomor 361/M/KEP/2025 menetapkan indikator dampak "
          "sosial, ekonomi, dan lingkungan perguruan tinggi yang dikelompokkan ke dalam 14 tema. Laporan ini memetakan "
          "jejak kegiatan Universitas Gadjah Mada (UGM) yang dipublikasikan ke tema-tema tersebut"
          + (", serta keterkaitannya dengan Tujuan Pembangunan Berkelanjutan (Sustainable Development Goals/SDGs)."
             if mode != "impact" else "."))
    lap.p("Laporan disusun secara otomatis oleh UGM Analytics dari data yang tersedia pada saat pembuatan, sebagai "
          "bahan awal yang membantu unit penyusun menelusuri bukti kegiatan per tema sebelum pengukuran indikator "
          "resmi dilakukan.")


def _tujuan(lap: _Laporan, mode: str) -> None:
    lap.p("Tujuan penyusunan laporan ini adalah:")
    items = [
        "Memetakan jejak kegiatan UGM yang dipublikasikan ke tema dampak sosial, ekonomi, dan lingkungan Kepmen 361/M/KEP/2025.",
        "Menunjukkan tema, periode, dan fakultas/unit kerja yang paling banyak mendokumentasikan kegiatan berdampak.",
        "Menyediakan daftar bukti awal (berita dan mata kuliah) yang dapat ditelusuri unit penyusun laporan resmi.",
    ]
    if mode != "impact":
        items.insert(2, "Menggambarkan keterkaitan kegiatan UGM dengan 17 Tujuan Pembangunan Berkelanjutan (SDGs).")
    lap.daftar(items, bernomor=True)


def _unit_analisis(lap: _Laporan, story: dict[str, Any], sumber: dict[str, Any] | None) -> None:
    b = _berita_sumber(sumber)
    mk = (sumber or {}).get("mata_kuliah") or {}
    lap.p("Unit analisis laporan ini adalah berita yang dipublikasikan pada portal resmi UGM dan mata kuliah yang "
          "ditawarkan program studi. Setiap berita dipetakan ke tema Kepmen melalui pencocokan kata kunci pada judul, "
          "deskripsi, dan isi berita; fakultas/unit kerja dikenali dari penyebutan nama resmi 44 fakultas, sekolah, dan "
          "unit kerja UGM di dalam berita.")
    baris = []
    if b:
        baris.append(["Berita portal " + str(b.get("situs", "ugm.ac.id")),
                      f"{ang(b.get('diambil', 0))} berita diambil dari {ang(b.get('sitemap', 0))} URL sitemap; "
                      f"{ang(b.get('berdampak', 0))} memuat konten dampak",
                      f"{b.get('tahun_awal', '')}–{b.get('tahun_akhir', '')}"])
    if mk.get("tersedia"):
        baris.append(["Mata kuliah (web program studi)",
                      f"{ang(mk.get('baris', 0))} baris penawaran dari {ang(mk.get('prodi', 0))} program studi dan "
                      f"{ang(mk.get('fakultas', 0))} fakultas/sekolah; {ang(mk.get('mk_unik', 0))} MK unik",
                      "Kurikulum berjalan"])
    if baris:
        lap.tabel("Sumber data laporan", ["Sumber", "Cakupan", "Periode"], baris,
                  sumber="Ringkasan sumber data UGM Analytics")
    lap.p("Analisis bersifat deskriptif. Angka yang disajikan adalah jumlah berita unik (satu berita dihitung sekali "
          "per tema walaupun memuat beberapa kata kunci) dan bersifat batas bawah: kegiatan yang tidak diberitakan tidak "
          "tercakup. Angka ini tidak menggantikan nilai capaian indikator resmi yang dihitung dari data administratif "
          "unit kerja sesuai kriteria Kepmen.")


def _identifikasi(lap: _Laporan, story: dict[str, Any], filter_teks: list[str], hari_ini: date) -> None:
    lap.bab("LEMBAR IDENTIFIKASI")
    rows = [
        ["Nama Perguruan Tinggi", "Universitas Gadjah Mada"],
        ["Status", "Perguruan Tinggi Negeri Badan Hukum"],
        ["Jenis laporan", f"Laporan dampak berbasis pemberitaan dan kurikulum (analisis {MODE_JUDUL[story['mode']]})"],
        ["Periode data", f"{story['filters']['year_from']}–{story['filters']['year_to']}"],
        ["Filter", "; ".join(filter_teks) or "Semua data"],
        ["Data per", str(story.get("data_as_of") or "-")],
        ["Tanggal penyusunan", tanggal_id(hari_ini)],
        ["Disusun dengan", "UGM Analytics (otomatis)"],
        ["Koordinator dan pengesahan", "Diisi oleh unit penyusun sebelum laporan disahkan"],
    ]
    lap.tabel("Identitas laporan", ["Butir", "Keterangan"], rows, kunci_nilai=True)


def _filter_teks(story: dict[str, Any], label_tema: dict[str, str]) -> list[str]:
    f = story["filters"]
    teks = []
    if f.get("pillars") and len(f["pillars"]) < 3 and story["mode"] != "sdgs":
        teks.append("Dampak: " + ", ".join(f["pillars"]))
    if f.get("topics") and len(f["topics"]) < len(label_tema):
        teks.append("Tema: " + ", ".join(label_tema.get(t, t) for t in f["topics"]))
    if f.get("sdgs"):
        teks.append("SDG: " + ", ".join(str(s) for s in f["sdgs"]))
    if f.get("units"):
        teks.append("Fakultas/unit kerja: " + ", ".join(f["units"]))
    return teks


# ------------------------------------------------------------------ mode dampak
def _dasar_mk(mk: dict[str, Any]) -> str:
    dasar = str(mk.get("dasar", ""))
    if dasar.startswith("Indikator resmi"):
        return "Angka ini merupakan indikator resmi Kepmen untuk kurikulum (mata kuliah berstatus substansial)."
    if dasar.startswith("Kriteria"):
        return ("Keterkaitan berasal dari kriteria kurasi manual a–j pada mata kuliah berstatus substansial; "
                "ini perluasan analitik, bukan indikator tema ini.")
    if dasar.startswith("Keyword"):
        return ("Keterkaitan ditentukan dari kata kunci pada nama dan deskripsi mata kuliah, sehingga menunjukkan "
                "kedekatan topik, bukan indikator tema ini.")
    return _kalimat(mk.get("catatan"))


def _sub_tema(lap: _Laporan, s: dict[str, Any], mode: str, start: str, end: str) -> None:
    judul = s.get("report_title") or s["label"]
    lap.sub(f"{s['number']} {judul}".strip())
    n = int(s["metrics"][0]["value"]) if s.get("metrics") else 0
    rentang = s["metrics"][1]["value"] if len(s.get("metrics", [])) > 1 else "-"
    indikator = (s.get("indicator") or "").strip()
    if indikator:
        lap.p(f"Dalam Kepmen 361/M/KEP/2025, tema {judul} diukur dengan indikator "
              f"{indikator[0].lower() + indikator[1:]}{'' if indikator.endswith('.') else '.'} "
              f"Rincian indikator resmi disajikan pada tabel berikut.")
        rows = [["Indikator", indikator], ["Definisi", s.get("definition", "")], ["Kriteria", s.get("criteria", "")],
                ["Formula", s.get("formula", "")], ["Satuan", s.get("unit", "")],
                ["Klaster SDG", ", ".join(x["label"] for x in s.get("sdg_labels", [])) or "-"]]
        lap.tabel(f"Indikator resmi tema {judul}", ["Butir", "Uraian"], [r for r in rows if r[1]],
                  sumber="Kepmen 361/M/KEP/2025", kunci_nilai=True)

    charts = s.get("charts", [])
    if n == 0:
        lap.p(f"Pada periode {start}–{end} dan filter yang dipilih, belum ditemukan berita UGM yang memuat kata kunci "
              f"tema {judul}. Bukti capaian tema ini perlu dihimpun langsung dari unit kerja terkait.")
    else:
        tren = _chart(charts, "_tren")
        lap.p(f"Selama periode {start}–{end}, portal berita UGM memuat {ang(n)} berita yang berkaitan dengan tema "
              f"{judul}, tersebar pada rentang tahun {rentang}. " + _kalimat(tren.get("insight") if tren else ""))
        if tren:
            pts = tren["data"]["series"][0]["points"]
            lap.gambar(gambar_garis([p["x"] for p in pts], [p["y"] for p in pts]),
                       f"Tren pemberitaan tema {judul} per tahun", "Portal berita UGM, diolah UGM Analytics")

        unit = _chart(charts, "_unit")
        if unit and unit.get("data"):
            top = unit["data"][:10]
            lap.p(f"Berdasarkan penyebutan fakultas/unit kerja dalam berita, penyumbang pemberitaan terbesar tema ini adalah "
                  + _daftar_kata([f"{d['label']} ({ang(d['value'])} berita)" for d in top[:3]]) + ". "
                  "Angka ini menunjukkan unit yang paling sering disebut, bukan besaran kontribusi final tiap unit.")
            lap.gambar(gambar_batang([d["label"] for d in top], [d["value"] for d in top], horizontal=True),
                       f"Fakultas/unit kerja penyumbang berita tema {judul}", "Portal berita UGM, diolah UGM Analytics")

        sdg = _chart(charts, "_sdg")
        if mode == "impact-sdgs" and sdg and sdg.get("data"):
            data = sorted(sdg["data"], key=lambda d: int(str(d["label"]).split()[-1]))
            lap.p(f"Dilihat dari Tujuan Pembangunan Berkelanjutan, berita tema ini terhubung dengan "
                  f"{len(data)} SDG. " + _kalimat(sdg.get("insight")))
            lap.gambar(gambar_batang([d["label"] for d in data], [d["value"] for d in data]),
                       f"Sebaran berita tema {judul} per SDG", "Pemetaan resmi tema Kepmen ke klaster SDG")
        elif s.get("sdg_labels"):
            lap.p("Tema ini termasuk klaster " + _daftar_kata([x["label"] for x in s["sdg_labels"]]) + ".")

    mk = s.get("mata_kuliah") or {}
    if mk.get("jumlah"):
        nama = [r.get("nama_mk", "") for r in (mk.get("tabel") or {}).get("rows", [])[:5]]
        lap.p(f"Dari sisi kurikulum, terdapat {ang(mk['jumlah'])} mata kuliah dari {ang(mk.get('fakultas', 0))} "
              f"fakultas/sekolah yang terkait dengan tema ini"
              + (f", antara lain {_daftar_kata(nama)}" if nama else "") + ". " + _dasar_mk(mk))
    elif mk.get("catatan"):
        lap.p("Dari sisi kurikulum, belum ada mata kuliah yang dipetakan ke tema ini. " + _kalimat(mk["catatan"]))

    tabel = (s.get("tables") or [{}])[0]
    rows = (tabel.get("rows") or [])[:5]
    if rows:
        lap.tabel(f"Contoh pemberitaan terbaru tema {judul}", ["Tanggal", "Judul", "Tautan"],
                  [[r.get("tanggal", ""), r.get("judul", ""), r.get("url", "")] for r in rows], links={1: 2},
                  sumber="Portal berita UGM")
    if n:
        satuan = s.get("unit") or "satuan indikator"
        lap.p(f"Secara keseluruhan, tema {judul} tercermin dalam {ang(n)} berita UGM selama periode {start}–{end}. "
              f"Angka tersebut menggambarkan jejak kegiatan yang dipublikasikan, bukan nilai capaian indikator "
              f"dalam satuan {satuan.lower()}; nilai resmi indikator tetap memerlukan data unit kerja sesuai "
              f"kriteria Kepmen.")


def _bab_dampak(lap: _Laporan, ch: dict[str, Any], mode: str, start: str, end: str) -> None:
    pilar = ch["pillar"]
    lap.bab(f"{ch['chapter']} - {ch['title'].upper()}")
    lap.p(PENGANTAR_PILAR.get(pilar, ""))
    subs = ch.get("subsections", [])
    nama_tema = [s.get("report_title") or s["label"] for s in subs]
    dist = _chart(ch.get("charts", []), "_tema")
    lap.p(f"Pada periode {start}–{end}, tercatat {ang(ch.get('total', 0))} berita UGM bertema dampak {pilar.lower()} "
          f"yang tersebar pada {len(subs)} tema, yaitu {_daftar_kata(nama_tema)}. "
          + _kalimat(dist.get("insight") if dist else "")
          + " Satu berita dapat masuk ke lebih dari satu tema.")
    if dist and dist.get("data"):
        lap.gambar(gambar_batang([d["label"] for d in dist["data"]], [d["value"] for d in dist["data"]],
                                 horizontal=True, warna=PILAR_WARNA.get(pilar, NAVY)),
                   f"Distribusi berita per tema dampak {pilar.lower()}", "Portal berita UGM, diolah UGM Analytics")
    heat = _chart(ch.get("charts", []), "_tema_tahun")
    if heat and heat.get("data", {}).get("rows"):
        d = heat["data"]
        lap.p("Sebaran pemberitaan tiap tema per tahun disajikan pada gambar berikut; warna yang lebih pekat "
              "menandakan berita yang lebih banyak. " + _kalimat(heat.get("insight")))
        lap.gambar(gambar_heatmap(d["rows"], d["cols"], d["values"]),
                   f"Tema × tahun dampak {pilar.lower()}", "Portal berita UGM, diolah UGM Analytics")
    for s in subs:
        _sub_tema(lap, s, mode, start, end)


def _laporan_dampak(lap: _Laporan, story: dict[str, Any], sumber: dict[str, Any] | None) -> None:
    mode, f = story["mode"], story["filters"]
    start, end = f["year_from"], f["year_to"]
    chapters = story.get("chapters", [])
    overview = story.get("overview", [])
    total = next((m["value"] for m in story["executive"]["metrics"] if m["label"] == "Total berita dampak"), 0)

    # BAB I
    lap.bab("BAB I - PENDAHULUAN")
    lap.sub("1.1 Latar Belakang")
    _latar_belakang(lap, mode)
    lap.sub("1.2 Tujuan")
    _tujuan(lap, mode)
    lap.sub("1.3 Ruang Lingkup")
    lap.p(f"Ruang lingkup laporan mencakup pemberitaan UGM periode {start}–{end} yang memuat konten dampak, "
          f"dikelompokkan ke dalam aspek berikut:")
    for i, ch in enumerate(chapters):
        lap.sub(f"{HURUF[i]}. Dampak {ch['pillar']}", level=3, masuk_daftar=False)
        lap.p(PENGANTAR_PILAR.get(ch["pillar"], ""))
    lap.sub("1.4 Unit Analisis")
    _unit_analisis(lap, story, sumber)

    for ch in chapters:
        _bab_dampak(lap, ch, mode, start, end)

    # BAB V
    # Nomor bab mengikuti daftar isi resmi (Lingkungan tetap BAB IV walau hanya pilar itu yang dipilih),
    # supaya sub-bab 4.x tetap cocok dengan laporan UGM 2025; kesimpulan selalu BAB V.
    lap.bab("BAB V - KESIMPULAN")
    if overview:
        urut = sorted(overview, key=lambda o: -o["total"])
        lap.p(f"Selama periode {start}–{end}, kegiatan UGM yang berdampak tercermin dalam {ang(total)} berita unik. "
              f"Menurut aspek, pemberitaan terbanyak ada pada dampak "
              + _daftar_kata([f"{o['pillar'].lower()} ({ang(o['total'])} berita)" for o in urut]) + ".")
        lap.p("Tema yang paling banyak diberitakan pada tiap aspek adalah "
              + _daftar_kata([f"{o['top_topic']} untuk dampak {o['pillar'].lower()} ({ang(o['top_topic_count'])} berita)"
                              for o in urut if o.get("top_topic")]) + ".")
    semua = [s for ch in chapters for s in ch.get("subsections", [])]
    sedikit = sorted(semua, key=lambda s: s["metrics"][0]["value"])[:3]
    if sedikit:
        lap.p("Tema dengan pemberitaan paling sedikit adalah "
              + _daftar_kata([f"{s.get('report_title') or s['label']} ({ang(s['metrics'][0]['value'])} berita)"
                              for s in sedikit])
              + ". Tema-tema ini perlu mendapat perhatian dalam pengumpulan bukti, baik melalui dokumentasi kegiatan "
                "yang lebih rutin maupun data administratif unit kerja.")
    mk = story.get("mata_kuliah") or {}
    resmi = next((m["value"] for m in mk.get("metrics", []) if "indikator resmi" in m["label"]), None)
    if mk.get("tersedia") and resmi is not None:
        lap.p(f"Dari sisi kurikulum, {ang(resmi)} mata kuliah unik berstatus substansial memenuhi indikator resmi "
              f"tema Pendidikan dan Penelitian (4.5).")
    lap.p("Untuk penyusunan laporan periode berikutnya, pemetaan berbasis pemberitaan ini dapat dipakai sebagai "
          "penunjuk awal bukti per tema, sementara nilai capaian indikator resmi dihimpun dari sistem informasi "
          "dan data administratif unit kerja agar hasilnya akurat, lengkap, dan mutakhir.")


# ------------------------------------------------------------------ mode SDGs
def _laporan_sdgs(lap: _Laporan, story: dict[str, Any], sumber: dict[str, Any] | None) -> None:
    f = story["filters"]
    start, end = f["year_from"], f["year_to"]
    m = {x["label"]: x["value"] for x in story["executive"]["metrics"]}
    charts = story.get("cross", {}).get("charts", [])
    peta = story.get("sdg_peta") or {}
    tiles = [t for t in peta.get("tiles", []) if not f.get("sdgs") or t["sdg"] in f["sdgs"]]
    n_tag = int(m.get("Berita bertanda SDG", 0) or 0)

    lap.bab("BAB I - PENDAHULUAN")
    lap.sub("1.1 Latar Belakang")
    _latar_belakang(lap, "sdgs")
    lap.sub("1.2 Tujuan")
    _tujuan(lap, "sdgs")
    lap.sub("1.3 Ruang Lingkup")
    lap.p(f"Ruang lingkup laporan mencakup seluruh URL berita pada sitemap portal UGM periode {start}–{end} "
          f"({ang(m.get('Total berita (sitemap)', 0))} berita) yang dipetakan ke 17 SDGs. Satu berita dapat terhubung "
          f"dengan lebih dari satu SDG.")
    lap.sub("1.4 Unit Analisis")
    _unit_analisis(lap, story, sumber)
    lap.p("Pemetaan SDG memakai kamus kata kunci per SDG yang dicocokkan pada slug URL, judul, deskripsi, dan isi "
          "berita, ditambah tanda SDG yang diberikan manual oleh tim untuk berita yang tidak memuat kata kunci.")

    lap.bab("BAB II - SEBARAN PEMBERITAAN PER SDG")
    lap.p(id_angka(story["executive"].get("narrative", "")))
    lap.sub("2.1 Ringkasan per SDG")
    ringkas = _chart(charts, "ringkasan_sdg")
    if ringkas and ringkas.get("data"):
        lap.p(_kalimat(ringkas.get("insight")) + " Gambar berikut menyajikan jumlah berita untuk setiap SDG, "
              "diurutkan menurut nomor SDG.")
        lap.gambar(gambar_batang([d["label"] for d in ringkas["data"]], [d["value"] for d in ringkas["data"]]),
                   "Jumlah berita per SDG (urut SDG 1–17)", "Portal berita UGM, diolah UGM Analytics")
        lap.tabel("Jumlah berita per SDG", ["SDG", "Nama", "Jumlah berita", "Porsi"],
                  [[d["label"], d.get("detail") or "", ang(d["value"]), pct(d["value"], n_tag)] for d in ringkas["data"]],
                  sumber="Porsi dihitung terhadap berita bertanda SDG; satu berita dapat masuk beberapa SDG")
    heat = _chart(charts, "sdg_tahun_heatmap")
    if heat and heat.get("data", {}).get("rows"):
        lap.sub("2.2 Perkembangan per Tahun")
        d = heat["data"]
        lap.p(_kalimat(heat.get("insight")) + " Gambar berikut memperlihatkan intensitas pemberitaan setiap SDG per tahun.")
        lap.gambar(gambar_heatmap(d["rows"], d["cols"], d["values"]), "Jumlah berita per SDG dan tahun",
                   "Portal berita UGM, diolah UGM Analytics")
    unit = _chart(charts, "unit")
    if unit and unit.get("data"):
        lap.sub("2.3 Kontribusi Fakultas dan Unit Kerja")
        top = unit["data"][:12]
        lap.p(_kalimat(unit.get("insight")) + " Unit dikenali dari penyebutan nama resmi dalam berita, sehingga angka ini "
              "bersifat batas bawah.")
        lap.gambar(gambar_batang([d["label"] for d in top], [d["value"] for d in top], horizontal=True),
                   "Fakultas/unit kerja dengan berita bertanda SDG terbanyak", "Portal berita UGM, diolah UGM Analytics")

    lap.bab("BAB III - PROFIL PER SDG")
    lap.p("Subbab berikut merangkum setiap SDG: jumlah berita, porsinya terhadap seluruh berita bertanda SDG, dan "
          "kata kunci yang paling sering muncul pada berita tersebut.")
    mk_chart = _chart((story.get("mata_kuliah") or {}).get("charts", []), "matkul_sdg_langsung")
    mk_per_sdg = {d["label"]: d["value"] for d in (mk_chart or {}).get("data", [])}
    for i, t in enumerate(tiles, start=1):
        lap.sub(f"3.{i} SDG {t['sdg']}: {t['nama']}")
        if not t["jumlah"]:
            lap.p(f"Pada periode dan filter ini belum ditemukan berita UGM yang terhubung dengan SDG {t['sdg']}.")
            continue
        kw = [k for k in t.get("keywords", []) if k.get("jumlah")][:5]
        teks = (f"SDG {t['sdg']} ({t['nama']}) terhubung dengan {ang(t['jumlah'])} berita UGM, atau "
                f"{pct(t['jumlah'], n_tag)} dari seluruh berita bertanda SDG.")
        if kw:
            teks += " Kata kunci yang paling sering muncul adalah " + _daftar_kata(
                [f"“{k['keyword']}” ({ang(k['jumlah'])} berita)" for k in kw]) + "."
        mk_n = mk_per_sdg.get(f"SDG {t['sdg']}")
        if mk_n:
            teks += f" Di kurikulum, {ang(mk_n)} mata kuliah memuat kata kunci SDG ini."
        lap.p(teks)

    mk = story.get("mata_kuliah") or {}
    if mk.get("tersedia") and mk.get("metrics"):
        lap.bab("BAB IV - KURIKULUM DAN SDGs")
        mm = {x["label"]: x["value"] for x in mk["metrics"]}
        lap.p(f"Sebanyak {ang(mm.get('MK unik ter-tag SDG', 0))} mata kuliah unik memuat kata kunci minimal satu SDG, "
              f"ditawarkan oleh {ang(mm.get('Fakultas/sekolah terlibat', 0))} fakultas/sekolah. "
              + _kalimat(mk_chart.get("insight") if mk_chart else "")
              + " Kamus kata kunci yang dipakai sama dengan kamus berita, sehingga angka ini indikatif dan bukan "
                "indikator Kepmen.")
        if mk_chart and mk_chart.get("data"):
            lap.gambar(gambar_batang([d["label"] for d in mk_chart["data"]], [d["value"] for d in mk_chart["data"]],
                                     satuan="MK"),
                       "Mata kuliah per SDG", "Web program studi UGM, diolah UGM Analytics")
        fak = _chart(mk.get("charts", []), "matkul_fakultas")
        if fak and fak.get("data"):
            lap.gambar(gambar_batang([d["label"] for d in fak["data"]], [d["value"] for d in fak["data"]],
                                     horizontal=True, satuan="MK"),
                       "Fakultas/sekolah dengan mata kuliah ter-tag SDG terbanyak",
                       "Web program studi UGM, diolah UGM Analytics")
        bab_akhir = "V"
    else:
        bab_akhir = "IV"

    lap.bab(f"BAB {bab_akhir} - KESIMPULAN")
    lap.p(f"Selama periode {start}–{end}, {ang(n_tag)} dari {ang(m.get('Total berita (sitemap)', 0))} berita UGM "
          f"({m.get('Cakupan', '-')}) terhubung dengan minimal satu SDG.")
    ada = sorted([t for t in tiles if t["jumlah"]], key=lambda t: -t["jumlah"])
    if ada:
        lap.p("SDG yang paling banyak disentuh adalah "
              + _daftar_kata([f"SDG {t['sdg']} {t['nama']} ({ang(t['jumlah'])} berita)" for t in ada[:3]])
              + ", sedangkan yang paling sedikit adalah "
              + _daftar_kata([f"SDG {t['sdg']} {t['nama']} ({ang(t['jumlah'])} berita)" for t in ada[-3:][::-1]])
              + ".")
    if story.get("tanpa_sdg_total"):
        lap.p(f"Masih terdapat {ang(story['tanpa_sdg_total'])} berita yang belum bertanda SDG; berita ini dapat ditandai "
              "manual melalui dashboard agar pemetaan lebih lengkap.")
    lap.p("Pemetaan berbasis kata kunci ini dapat dipakai sebagai penunjuk awal kontribusi UGM terhadap SDGs. Untuk "
          "pelaporan resmi, klaim per SDG perlu diperkuat dengan data kegiatan dan capaian dari unit kerja terkait.")


# ------------------------------------------------------------------ titik masuk
def build_laporan(story: dict[str, Any], sumber: dict[str, Any] | None = None,
                  label_tema: dict[str, str] | None = None, hari_ini: date | None = None) -> dict[str, Any]:
    """Susun model dokumen laporan dari payload story (filter aktif) + ringkasan sumber."""
    hari_ini = hari_ini or date.today()
    mode = story["mode"]
    f = story["filters"]
    start, end = f["year_from"], f["year_to"]
    lap = _Laporan()
    filter_teks = _filter_teks(story, label_tema or {})
    judul = ("LAPORAN KONTRIBUSI UGM TERHADAP TUJUAN PEMBANGUNAN BERKELANJUTAN (SDGs)" if mode == "sdgs"
             else "LAPORAN DAMPAK SOSIAL, EKONOMI, DAN LINGKUNGAN")
    sub = "Berbasis pemberitaan portal UGM dan kurikulum program studi"

    # Ringkasan eksekutif
    lap.bab("RINGKASAN EKSEKUTIF")
    exe = story.get("executive", {})
    if mode == "sdgs":
        lap.p("Laporan ini memetakan pemberitaan kegiatan Universitas Gadjah Mada ke 17 Tujuan Pembangunan "
              "Berkelanjutan (SDGs) untuk menunjukkan bidang yang paling banyak disentuh dan bidang yang masih perlu "
              "diperkuat dokumentasinya.")
        lap.p(id_angka(exe.get("narrative", "")))
    else:
        lap.p("Laporan ini memetakan kegiatan Universitas Gadjah Mada yang dipublikasikan ke tema dampak sosial, "
              "ekonomi, dan lingkungan sesuai Keputusan Menteri Pendidikan Tinggi, Sains, dan Teknologi Nomor "
              "361/M/KEP/2025" + (", beserta keterkaitannya dengan SDGs" if mode == "impact-sdgs" else "") + ".")
        lap.p(id_angka(exe.get("narrative", "")))
        ov = story.get("overview", [])
        if ov:
            lap.gambar(gambar_batang([o["pillar"] for o in ov], [o["total"] for o in ov],
                                     warna=[PILAR_WARNA.get(o["pillar"], NAVY) for o in ov]),
                       f"Jumlah berita per aspek dampak, {start}–{end}", "Portal berita UGM, diolah UGM Analytics")
    b = _berita_sumber(sumber)
    if b.get("diambil"):
        lap.p(f"Data dasar berasal dari {ang(b['diambil'])} berita portal {b.get('situs', 'ugm.ac.id')} periode "
              f"{b.get('tahun_awal')}–{b.get('tahun_akhir')}, dengan {ang(b.get('berdampak', 0))} berita "
              f"({pct(b.get('berdampak', 0), b['diambil'])}) memuat konten dampak. Seluruh angka dalam laporan ini "
              f"mengikuti filter yang dipilih saat laporan dibuat.")

    _identifikasi(lap, story, filter_teks, hari_ini)
    if mode == "sdgs":
        _laporan_sdgs(lap, story, sumber)
    else:
        _laporan_dampak(lap, story, sumber)

    lap.bab("REFERENSI")
    lap.daftar([
        "Kementerian Pendidikan Tinggi, Sains, dan Teknologi. (2025). Keputusan Menteri Pendidikan Tinggi, Sains, dan "
        "Teknologi Nomor 361/M/KEP/2025 tentang Indikator Dampak Sosial, Ekonomi, dan Lingkungan Perguruan Tinggi. "
        "Jakarta.",
        "Universitas Gadjah Mada. (2025). Laporan Dampak Sosial, Ekonomi, dan Lingkungan UGM 2025. Yogyakarta.",
        f"Universitas Gadjah Mada. Portal berita ugm.ac.id, diakses melalui sitemap (data per {story.get('data_as_of') or '-'}).",
        "Universitas Gadjah Mada. Deskripsi Matkul Kepmen dan Ringkasan Indikator Kepmen (kurasi kurikulum program studi).",
    ])
    lap.bab("LAMPIRAN")
    lap.sub("Catatan metodologi", masuk_daftar=True)
    lap.daftar(list(story.get("caveats", [])) + list((story.get("mata_kuliah") or {}).get("catatan_metode", [])))

    return {
        "title": judul,
        "subtitle": sub,
        "institution": "UNIVERSITAS GADJAH MADA",
        "period": f"{start}–{end}",
        "mode": mode,
        "mode_label": MODE_JUDUL[mode],
        "filters": filter_teks,
        "generated": tanggal_id(hari_ini),
        "data_as_of": story.get("data_as_of"),
        "toc": lap.daftar_isi,
        "figures": lap.daftar_gambar,
        "tables": lap.daftar_tabel,
        "blocks": lap.blocks,
    }


# ------------------------------------------------------------------ suntingan pratinjau
MAX_SUNTINGAN = 500
MAX_TEKS_PARAGRAF = 5000


def terapkan_suntingan(lap: dict[str, Any], suntingan: dict[str, str] | None) -> dict[str, Any]:
    """Salinan laporan dengan teks paragraf yang disunting di pratinjau.

    Hanya blok `paragraph` yang boleh diubah (judul, gambar, tabel, dan angka tetap dari data).
    Paragraf yang dikosongkan dihapus dari dokumen. Laporan di cache tidak diubah.
    """
    if not suntingan:
        return lap
    if len(suntingan) > MAX_SUNTINGAN:
        raise ValueError(f"Maksimal {MAX_SUNTINGAN} paragraf disunting per unduhan.")
    blocks = list(lap["blocks"])
    for kunci, teks in suntingan.items():
        try:
            i = int(kunci)
        except (TypeError, ValueError) as exc:
            raise ValueError("Indeks suntingan tidak valid.") from exc
        if not 0 <= i < len(blocks) or blocks[i]["type"] != "paragraph":
            raise ValueError("Hanya paragraf narasi yang bisa disunting.")
        teks = str(teks or "").strip()
        if len(teks) > MAX_TEKS_PARAGRAF:
            raise ValueError(f"Paragraf maksimal {MAX_TEKS_PARAGRAF} karakter.")
        blocks[i] = {**blocks[i], "text": teks}
    return {**lap, "blocks": [b for b in blocks if not (b["type"] == "paragraph" and not b["text"])]}


# ------------------------------------------------------------------ render .docx
def _font(run: Any, size: float | None = None, bold: bool | None = None, color: str | None = None,
          italic: bool | None = None) -> None:
    run.font.name = "Arial"
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def _hyperlink(paragraph: Any, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                    is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0563C1"), ("w:u", "single"), ("w:sz", "16")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        props.append(el)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    props.append(fonts)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _shade(cell: Any, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tc_pr.append(shd)


def _nomor_halaman(section: Any) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    _font(run, 9, color="#555555")


def _caption(doc: Any, text: str, sumber: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run(text), 9.5, bold=True)
    if sumber:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(s.add_run(f"Sumber: {sumber}"), 8.5, italic=True, color="#555555")


def _daftar_ringkas(doc: Any, judul: str, items: list[str], indent: dict[int, int] | None = None) -> None:
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_after = Pt(12)
    _font(h.add_run(judul), 13, bold=True, color=NAVY)
    if not items:
        _font(doc.add_paragraph().add_run("Tidak ada."), 10.5)
    for i, text in enumerate(items):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.8 * (indent or {}).get(i, 0))
        _font(p.add_run(text), 10.5, bold=(indent or {}).get(i, 0) == 0 and indent is not None)


def render_docx(lap: dict[str, Any]) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    # Sampul
    for _ in range(6):
        doc.add_paragraph()
    for text, size, color in ((lap["title"], 20, NAVY), (lap["subtitle"], 12, "#33414e"),
                              (f"Periode data {lap['period']}", 12, "#33414e")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        _font(p.add_run(text), size, bold=size > 12, color=color)
    for _ in range(10):
        doc.add_paragraph()
    for text, size in ((lap["institution"], 14), ("Yogyakarta", 11), (lap["generated"], 11)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(text), size, bold=size > 11, color=NAVY)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    _nomor_halaman(body)
    blocks = lap["blocks"]
    toc_disisip = False
    bab_ke = 0
    baru_daftar = False
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            # Daftar isi/gambar/tabel disisipkan tepat sebelum BAB I, seperti urutan laporan resmi.
            if block["level"] == 1 and block["text"].startswith("BAB I ") and not toc_disisip:
                toc_disisip = True
                for judul, items, indent in (
                    ("DAFTAR ISI", [e["text"] for e in lap["toc"]],
                     {i: e["level"] - 1 for i, e in enumerate(lap["toc"])}),
                    ("DAFTAR GAMBAR", lap["figures"], None),
                    ("DAFTAR TABEL", lap["tables"], None),
                ):
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                    _daftar_ringkas(doc, judul, items, indent)
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                baru_daftar = True
            if block["level"] == 1:
                # Bab pertama sudah di halaman baru (section setelah sampul); bab lain mulai halaman baru.
                if block.get("break_before") and bab_ke > 0 and not baru_daftar:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                bab_ke += 1
                baru_daftar = False
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(14)
                _font(p.add_run(block["text"]), 13, bold=True, color=NAVY)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12 if block["level"] == 2 else 8)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                _font(p.add_run(block["text"]), 12 if block["level"] == 2 else 11, bold=True,
                      color=NAVY if block["level"] == 2 else "#1f2d3a")
        elif kind == "paragraph":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(8)
            _font(p.add_run(block["text"]), 11)
        elif kind == "list":
            style = "List Number" if block.get("ordered") else "List Bullet"
            for item in block["items"]:
                p = doc.add_paragraph(style=style)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _font(p.add_run(item), 11)
        elif kind == "figure":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(BytesIO(base64.b64decode(block["image"])), width=Cm(15))
            _caption(doc, block["caption"], block.get("source"))
        elif kind == "table":
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.keep_with_next = True
            _font(cap.add_run(block["caption"]), 9.5, bold=True)
            cols = block["columns"]
            links = {int(k): v for k, v in (block.get("links") or {}).items()}
            link_cols = set(links.values())
            tampil = [i for i in range(len(cols)) if i not in link_cols]
            table = doc.add_table(rows=1, cols=len(tampil))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if not block.get("key_value"):
                for cell, i in zip(table.rows[0].cells, tampil):
                    _shade(cell, NAVY)
                    _font(cell.paragraphs[0].add_run(cols[i]), 9.5, bold=True, color="#ffffff")
            else:
                table._tbl.remove(table.rows[0]._tr)
            for row in block["rows"]:
                cells = table.add_row().cells
                for cell, i in zip(cells, tampil):
                    para = cell.paragraphs[0]
                    url = row[links[i]] if i in links else ""
                    if url.startswith(("http://", "https://")):
                        _hyperlink(para, url, row[i] or url)
                    else:
                        _font(para.add_run(row[i]), 9.5, bold=block.get("key_value") and i == 0)
            if block.get("source"):
                s = doc.add_paragraph()
                s.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _font(s.add_run(f"Sumber: {block['source']}"), 8.5, italic=True, color="#555555")
            else:
                doc.add_paragraph()
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
