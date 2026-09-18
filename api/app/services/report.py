from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor


NAVY = RGBColor(0x00, 0x21, 0x4A)


def _table(doc: Document, title: str, rows: list[dict[str, Any]]) -> None:
    doc.add_heading(title, level=1)
    if not rows:
        doc.add_paragraph("Tidak ada data untuk filter ini.")
        return
    columns = list(rows[0].keys())
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, column in zip(table.rows[0].cells, columns):
        cell.text = str(column)
    for row in rows[:40]:
        cells = table.add_row().cells
        for cell, column in zip(cells, columns):
            cell.text = str(row.get(column, ""))
    if len(rows) > 40:
        doc.add_paragraph(f"{len(rows) - 40} baris lain tidak ditampilkan.").italic = True


def build_report(mode: str, result: dict[str, Any]) -> bytes:
    document = Document()
    title = document.add_heading(f"Laporan Analisis {mode}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = NAVY
    document.add_paragraph("UGM Analytics — data live sesuai filter aktif.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Filter yang Diterapkan", level=1)
    for key, value in result.get("filters", {}).items():
        document.add_paragraph(f"{key}: {value}", style="List Bullet")
    document.add_heading("Ringkasan Eksekutif", level=1)
    document.add_paragraph(result.get("narrative", ""))
    document.add_heading("Metrik", level=1)
    metrics = [{"Metrik": key, "Nilai": value} for key, value in result.get("summary", {}).items()]
    _table(document, "Ringkasan", metrics)
    for name, rows in result.get("charts", {}).items():
        if isinstance(rows, list):
            _table(document, name.replace("_", " ").title(), rows)
    for name, rows in result.get("tables", {}).items():
        if isinstance(rows, list):
            _table(document, name.replace("_", " ").title(), rows)
    document.add_heading("Catatan Metodologi", level=1)
    for caveat in result.get("caveats", []):
        document.add_paragraph(caveat, style="List Bullet")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
