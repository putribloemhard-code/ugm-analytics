"""Template Excel per isian tabel LED/LKPS: unduh format kosong dan baca file yang diupload.

Acuan kolom = akreditasi/format_tabel.json (lewat registry: FORMAT_TABEL + kolom_dibutuhkan), jadi
template, form, dan upload selalu memakai kolom yang sama. Header bertingkat ("Grup > Sub") dirender
seperti file Excel resmi di akreditasi/docs/tabel_xlsx: judul tabel di baris 1, header ter-merge,
huruf Times New Roman, garis tipis.

Pengaman supaya data tidak masuk ke tabel yang salah: template membawa penanda isian di properti
berkas (keywords "ugm-analytics-tabel:<item_id>"); saat upload penanda itu, judul tabel, dan
seluruh teks header dicocokkan dengan format isian tujuan -- satu saja beda, file ditolak.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app.domain.source import load_accreditation_module

PENANDA = "ugm-analytics-tabel:"
BARIS_KOSONG = 30
MAX_BARIS = 2000
MAX_UKURAN = 5 * 1024 * 1024

FONT = "Times New Roman"
F_JUDUL = Font(name=FONT, size=12, bold=True)
F_HEADER = Font(name=FONT, size=11, bold=True)
F_ISI = Font(name=FONT, size=11)
TENGAH = Alignment(horizontal="center", vertical="center", wrap_text=True)
KIRI = Alignment(horizontal="left", vertical="top", wrap_text=True)
TIPIS = Side(style="thin", color="000000")
GARIS = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)


class TabelError(ValueError):
    """Isian tidak punya format tabel, atau file upload tidak cocok (400)."""


def _registry():
    return load_accreditation_module("registry_kebutuhan_data.py")


def _format(item_id: str) -> dict[str, Any]:
    fmt = _registry().FORMAT_TABEL.get(item_id)
    if fmt is None:
        raise TabelError("Isian ini tidak punya format tabel Excel.")
    return fmt


TANPA_NOMOR = {"Tanpa Nomor": " (tabel tanpa nomor)", "Tim Penyusun": ""}


def judul_tabel(fmt: dict[str, Any]) -> str:
    """'Tabel 1.A.1 Pimpinan ...'; tabel tanpa nomor resmi memakai judulnya saja."""
    kode = fmt["kode"]
    if vertikal(fmt):
        return f"{kode}. {fmt['judul']}"  # "A. Spesifikasi Program" (bagian Identitas Pengusul LED)
    if kode in TANPA_NOMOR:
        return fmt["judul"] + TANPA_NOMOR[kode]
    return f"Tabel {kode} {fmt['judul']}"


def vertikal(fmt: dict[str, Any]) -> bool:
    """Satu record: tiap butir satu baris (kolom A = butir, kolom B = keterangan)."""
    return fmt.get("bentuk") == "vertikal"


HEADER_VERTIKAL = ("Butir", "Keterangan")


def _jalur(fmt: dict[str, Any]) -> list[list[str]]:
    return [[b.strip() for b in k.split(">")] for k in fmt["kolom"]]


def _sel_header(fmt: dict[str, Any]) -> tuple[int, list[tuple[int, int, int, int, str]]]:
    """Kedalaman header + daftar sel (baris1, kolom1, baris2, kolom2, teks); baris/kolom 1-based,
    header mulai di baris 2 (baris 1 = judul)."""
    jalur = _jalur(fmt)
    dalam = max(len(j) for j in jalur)
    geser = 1 if fmt["nomor"] else 0
    sel: list[tuple[int, int, int, int, str]] = []
    if geser:
        sel.append((2, 1, 1 + dalam, 1, "No."))
    for lvl in range(dalam):
        i = 0
        while i < len(jalur):
            if len(jalur[i]) <= lvl:
                i += 1
                continue
            j = i
            while j + 1 < len(jalur) and len(jalur[j + 1]) > lvl and jalur[j + 1][:lvl + 1] == jalur[i][:lvl + 1]:
                j += 1
            daun = len(jalur[i]) == lvl + 1
            sel.append((2 + lvl, 1 + geser + i, (1 + dalam) if daun else (2 + lvl), 1 + geser + j, jalur[i][lvl]))
            i = j + 1
    return dalam, sel


def nama_file(fmt: dict[str, Any]) -> str:
    nama = re.sub(r'[\\/:*?"<>|]+', "-", f"Format {fmt['dokumen']} {judul_tabel(fmt)}")
    return f"{nama[:120].strip()}.xlsx"


def buat_template(item_id: str) -> tuple[bytes, str]:
    fmt = _format(item_id)
    if vertikal(fmt):
        return _template_vertikal(item_id, fmt)
    dalam, sel = _sel_header(fmt)
    n_kolom = len(fmt["kolom"]) + (1 if fmt["nomor"] else 0)
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "-", f"Tabel {fmt['kode']}")[:31]
    wb.properties.keywords = f"{PENANDA}{item_id}"
    wb.properties.title = judul_tabel(fmt)

    ws.cell(1, 1, judul_tabel(fmt)).font = F_JUDUL
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    if n_kolom > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_kolom)
    ws.row_dimensions[1].height = 22

    for r in range(2, 2 + dalam):
        for c in range(1, n_kolom + 1):
            ws.cell(r, c).border = GARIS
    for r1, c1, r2, c2, teks in sel:
        cell = ws.cell(r1, c1, teks)
        cell.font, cell.alignment = F_HEADER, TENGAH
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    awal = 2 + dalam
    for r in range(awal, awal + BARIS_KOSONG):
        for c in range(1, n_kolom + 1):
            cell = ws.cell(r, c)
            cell.border, cell.font, cell.alignment = GARIS, F_ISI, KIRI
        if fmt["nomor"]:
            ws.cell(r, 1, r - awal + 1).alignment = TENGAH

    lebar = [6] if fmt["nomor"] else []
    lebar += [min(max(len(j[-1]) + 4, 12), 40) for j in _jalur(fmt)]
    for i, w in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(awal, 1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), nama_file(fmt)


def _template_vertikal(item_id: str, fmt: dict[str, Any]) -> tuple[bytes, str]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Spesifikasi Program"
    wb.properties.keywords = f"{PENANDA}{item_id}"
    wb.properties.title = judul_tabel(fmt)
    ws.cell(1, 1, judul_tabel(fmt)).font = F_JUDUL
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws.row_dimensions[1].height = 22
    for c, teks in enumerate(HEADER_VERTIKAL, start=1):
        cell = ws.cell(2, c, teks)
        cell.font, cell.alignment, cell.border = F_HEADER, TENGAH, GARIS
    for r, butir in enumerate(fmt["kolom"], start=3):
        a, b = ws.cell(r, 1, butir), ws.cell(r, 2)
        a.font, a.alignment, a.border = F_ISI, KIRI, GARIS
        b.font, b.alignment, b.border = F_ISI, KIRI, GARIS
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 70
    ws.freeze_panes = "B3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), nama_file(fmt)


def _norm(teks: Any) -> str:
    return re.sub(r"\s+", " ", str(teks or "")).strip().lower()


def _nilai(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Ya" if v else "Tidak"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def baca_upload(item_id: str, isi: bytes) -> dict[str, Any]:
    """Baris data dari file template yang sudah diisi -> {"rows": [{kolom: nilai}], "kolom": [...]}.
    Tidak menyimpan apa pun: baris ditampilkan di form dulu, pengguna yang menekan Simpan."""
    fmt = _format(item_id)
    registry = _registry()
    if len(isi) > MAX_UKURAN:
        raise TabelError("Ukuran file melebihi 5 MB.")
    try:
        wb = load_workbook(io.BytesIO(isi), data_only=True)
    except Exception as exc:  # openpyxl melempar beragam galat untuk file rusak/bukan xlsx
        raise TabelError("File tidak bisa dibaca. Unggah file .xlsx dari tombol Unduh format.") from exc
    ws = wb.worksheets[0]
    judul = judul_tabel(fmt)

    penanda = str(wb.properties.keywords or "")
    if penanda.startswith(PENANDA):
        asal = penanda[len(PENANDA):].strip()
        if asal != item_id:
            asal_fmt = registry.FORMAT_TABEL.get(asal)
            nama_asal = judul_tabel(asal_fmt) if asal_fmt else asal
            raise TabelError(f"File ini format «{nama_asal}», bukan «{judul}». Unggah file dari tombol Unduh format di tabel ini.")
    if _norm(ws.cell(1, 1).value) != _norm(judul):
        raise TabelError(f"Judul di baris 1 file («{ws.cell(1, 1).value or 'kosong'}») tidak sama dengan «{judul}». "
                         "Pastikan file ini format tabel yang benar dan judulnya tidak diubah.")

    if vertikal(fmt):
        return _baca_vertikal(ws, fmt)
    dalam, sel = _sel_header(fmt)
    for r1, c1, _, _, teks in sel:
        ada = ws.cell(r1, c1).value
        if _norm(ada) != _norm(teks):
            raise TabelError(f"Header kolom {get_column_letter(c1)}{r1} di file berisi «{ada or 'kosong'}», "
                             f"seharusnya «{teks}». Jangan ubah header; unduh ulang format bila perlu.")

    kolom = [registry.kunci_kolom(k) for k in fmt["kolom"]]
    geser = 1 if fmt["nomor"] else 0
    rows: list[dict[str, str]] = []
    for nilai in ws.iter_rows(min_row=2 + dalam, max_col=geser + len(kolom), values_only=True):
        nilai = list(nilai) + [None] * (geser + len(kolom) - len(nilai))
        baris = {k: _nilai(v) for k, v in zip(kolom, nilai[geser:])}
        if all(v in ("", "-") for v in baris.values()):
            continue
        rows.append(baris)
        if len(rows) > MAX_BARIS:
            raise TabelError(f"File berisi lebih dari {MAX_BARIS} baris data.")
    if not rows:
        raise TabelError("Tidak ada baris data di file. Isi baris di bawah header lalu unggah lagi.")
    return {"kolom": kolom, "rows": rows}


def _baca_vertikal(ws, fmt: dict[str, Any]) -> dict[str, Any]:
    """Kolom A harus berisi butir persis urutan format; kolom B = nilainya. Hasil: satu record."""
    for c, teks in enumerate(HEADER_VERTIKAL, start=1):
        if _norm(ws.cell(2, c).value) != _norm(teks):
            raise TabelError(f"Header {get_column_letter(c)}2 di file berisi «{ws.cell(2, c).value or 'kosong'}», "
                             f"seharusnya «{teks}». Unduh ulang format bila perlu.")
    kolom = list(fmt["kolom"])
    baris: dict[str, str] = {}
    for i, butir in enumerate(kolom):
        ada = ws.cell(3 + i, 1).value
        if _norm(ada) != _norm(butir):
            raise TabelError(f"Sel A{3 + i} berisi «{ada or 'kosong'}», seharusnya butir «{butir}». "
                             "Jangan ubah, hapus, atau pindah baris butir.")
        baris[butir] = _nilai(ws.cell(3 + i, 2).value)
    if all(v in ("", "-") for v in baris.values()):
        raise TabelError("Kolom Keterangan masih kosong. Isi keterangan tiap butir lalu unggah lagi.")
    return {"kolom": kolom, "rows": [baris]}
