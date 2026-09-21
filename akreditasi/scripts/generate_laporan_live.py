"""Generator laporan LED & LKPS (.docx) -- Fase 3, sumber data MURNI live
(hasil pipeline Fase 2), BUKAN dokumen PDF lama.

Beda mendasar dari scripts/generate_template.py (generator LAMA):
- generate_template.py baca dari akreditasi_data_manual (sekarang tabel
  KOSONG, isi lama sudah diarsipkan ke akreditasi_data_manual_arsip_pdf --
  lihat scripts/migrasi_arsip_pdf.py). Skrip ini TIDAK menyentuh tabel
  arsip itu sama sekali.
- Status tiap item DIBACA dari akreditasi/data_source_map.json (Fase 1),
  BUKAN dari field status_ketersediaan di registry_kebutuhan_data.py
  (field itu sekarang cuma sisa kompatibilitas, lihat docstring modul
  registry poin 7 -- TIDAK dipakai sebagai acuan di sini).
- Isi utk item "tersedia" dibaca dari tabel-tabel hasil pipeline Fase 2:
  akreditasi_item_tersedia (5 item resmi), ditambah akreditasi_publikasi_dosen
  & akreditasi_berita_dcse sbg "Data pendukung (live)" tambahan pada item
  yang punya field sumber_tambahan di map (TETAP ditandai tidak_tersedia
  scr resmi -- data pendukung TIDAK mengubah status item).

Struktur dokumen (bab/urutan/kode item per Kriteria A-D utk LED, per
Bagian 1-6 utk LKPS) TETAP pakai registry_kebutuhan_data.py -- itu memang
struktur/daftar kebutuhan yang boleh diambil dari referensi PDF lama
(lihat instruksi user Fase 1), bukan isinya.

Placeholder 2 jenis (WAJIB beda, supaya pembaca tidak bingung, sesuai
instruksi user):
- jenis_kendala "akses_data"              -> "[DATA TIDAK TERSEDIA]" + catatan sumber
- jenis_kendala "perlu_penyusunan_manusia" -> "[NARASI PERLU DISUSUN TIM PENYUSUN PRODI]"
  (TANPA catatan soal sumber data -- itemnya memang di luar cakupan sistem ini)

Jalankan:
    ..\\venv\\Scripts\\python.exe scripts\\generate_laporan_live.py
"""
import json
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

import db  # noqa: E402
from registry_kebutuhan_data import (  # noqa: E402
    LABEL_BAGIAN_LKPS,
    LABEL_KRITERIA,
    URUTAN_BAGIAN_LKPS,
    URUTAN_KRITERIA,
    led_items_by_kriteria,
    lkps_items_by_bagian,
)

MAP_PATH = Path(__file__).resolve().parents[1] / "data_source_map.json"

WARNA_AKSES_DATA = RGBColor(0x99, 0x33, 0x33)       # merah tua -- soal akses/sistem
WARNA_PERLU_MANUSIA = RGBColor(0x1F, 0x4E, 0x79)    # biru tua -- soal penyusunan narasi
WARNA_SUMBER_TAMBAHAN = RGBColor(0x55, 0x55, 0x55)  # abu-abu -- info pendukung, bukan field resmi


def load_map() -> dict[str, dict]:
    with open(MAP_PATH, encoding="utf-8") as f:
        data = json.load(f)
    return {item["kode_item"]: item for item in data["items"]}


def load_item_tersedia(engine) -> pd.DataFrame:
    if not db.table_exists(engine, db.t("item_tersedia")):
        return pd.DataFrame(columns=["item_id", "baris_ke", "kolom", "nilai", "sumber_url"])
    return db.read_sql_retry(
        engine, f"SELECT item_id, baris_ke, kolom, nilai, sumber_url FROM `{db.t('item_tersedia')}`",
        label="load akreditasi_item_tersedia",
    )


def load_publikasi_dosen(engine) -> pd.DataFrame:
    if not db.table_exists(engine, db.t("publikasi_dosen")):
        return pd.DataFrame(columns=["dosen", "judul", "tahun", "sumber", "link"])
    return db.read_sql_retry(
        engine, f"SELECT dosen, judul, tahun, sumber, link FROM `{db.t('publikasi_dosen')}` "
                f"ORDER BY dosen, tahun DESC",
        label="load akreditasi_publikasi_dosen",
    )


def load_berita_dcse_relevan(engine) -> pd.DataFrame:
    if not db.table_exists(engine, db.t("berita_dcse")):
        return pd.DataFrame(columns=["judul", "tanggal", "url"])
    return db.read_sql_retry(
        engine, f"SELECT judul, tanggal, url FROM `{db.t('berita_dcse')}` WHERE relevan_mei=1 "
                f"ORDER BY tanggal DESC",
        label="load akreditasi_berita_dcse (relevan_mei)",
    )


def _rows_for_item(item_id: str, df: pd.DataFrame) -> list[dict]:
    """Pivot long (item_id, baris_ke, kolom, nilai) -> list of dict per baris_ke."""
    sub = df[df["item_id"] == item_id]
    if not len(sub):
        return []
    rows: dict[int, dict] = {}
    for _, r in sub.iterrows():
        rows.setdefault(int(r["baris_ke"]), {})[r["kolom"]] = r["nilai"] or ""
    return [rows[k] for k in sorted(rows)]


def _add_tersedia(doc: Document, item: dict, rows: list[dict]) -> None:
    """Render data live utk item "tersedia". PENTING (fix 2026-09-11 setelah
    review independen): kolom yang dirender WAJIB semua kolom_dibutuhkan
    resmi dari registry -- BUKAN cuma kolom yang kebetulan ada di
    akreditasi_item_tersedia. Sub-field resmi yang datanya belum berhasil
    ditarik pipeline ditandai EKSPLISIT "[DATA TIDAK TERSEDIA]" per sel/
    baris, BUKAN dihilangkan diam-diam (bug yang sudah terjadi: kolom
    "Pejabat Penandatangan" di identitas_pt_upps_ps dulu hilang tanpa
    keterangan apa pun). Kolom TAMBAHAN yang kebetulan ada di data tapi
    TIDAK ada di kolom_dibutuhkan resmi tetap ditampilkan di akhir (bukan
    dibuang) -- supaya tidak kehilangan data pendukung yang sudah ditarik."""
    kolom_resmi = list(item["kolom_dibutuhkan"])
    if not rows:
        p = doc.add_paragraph()
        p.add_run("[DATA TIDAK TERSEDIA]").bold = True
        p.runs[0].font.color.rgb = WARNA_AKSES_DATA
        doc.add_paragraph("Catatan: item ini berstatus \"tersedia\" di data_source_map.json tapi "
                           "pipeline belum berhasil menarik data pada run terakhir -- cek log pipeline.")
        return
    semua_kolom: list[str] = list(kolom_resmi)
    for row in rows:
        for k in row:
            if k not in semua_kolom:
                semua_kolom.append(k)  # kolom tambahan di luar kolom_dibutuhkan resmi -- tetap ditampilkan

    def _sel(row: dict, k: str) -> str:
        v = row.get(k, "")
        return str(v) if v else "[DATA TIDAK TERSEDIA]"

    if len(rows) == 1 and len(semua_kolom) <= 8:
        # narasi/identitas -- render sbg field: value, bukan tabel
        for k in semua_kolom:
            p = doc.add_paragraph()
            p.add_run(f"{k}: ").bold = True
            nilai = rows[0].get(k, "")
            if nilai:
                p.add_run(str(nilai))
            else:
                run = p.add_run("[DATA TIDAK TERSEDIA]")
                run.font.color.rgb = WARNA_AKSES_DATA
        return
    table = doc.add_table(rows=1, cols=len(semua_kolom))
    table.style = "Light Grid Accent 1"
    for i, k in enumerate(semua_kolom):
        table.rows[0].cells[i].text = k
        for r in table.rows[0].cells[i].paragraphs[0].runs:
            r.font.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, k in enumerate(semua_kolom):
            cells[i].text = _sel(row, k)


def _add_placeholder(doc: Document, map_entry: dict) -> None:
    jenis = map_entry.get("jenis_kendala")
    if jenis == "perlu_penyusunan_manusia":
        p = doc.add_paragraph()
        run = p.add_run("[NARASI PERLU DISUSUN TIM PENYUSUN PRODI]")
        run.bold = True
        run.font.color.rgb = WARNA_PERLU_MANUSIA
        return
    # default / akses_data
    p = doc.add_paragraph()
    run = p.add_run("[DATA TIDAK TERSEDIA]")
    run.bold = True
    run.font.color.rgb = WARNA_AKSES_DATA
    p2 = doc.add_paragraph()
    p2.add_run("Catatan: ").bold = True
    p2.add_run(map_entry.get("catatan_jika_tidak_tersedia", ""))


def _add_sumber_tambahan(doc: Document, map_entry: dict, kode_item: str,
                          df_pub: pd.DataFrame, df_berita: pd.DataFrame) -> None:
    teks = map_entry.get("sumber_tambahan")
    if not teks:
        return
    doc.add_heading("Data Pendukung (live, bukan pengganti field resmi di atas)", level=3)
    p = doc.add_paragraph()
    p.add_run(teks).italic = True
    p.runs[0].font.color.rgb = WARNA_SUMBER_TAMBAHAN

    if kode_item == "lkps_3_c_2" and len(df_pub):
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light List Accent 1"
        for i, k in enumerate(["Dosen", "Judul", "Tahun", "Sumber"]):
            table.rows[0].cells[i].text = k
        for _, r in df_pub.head(30).iterrows():
            cells = table.add_row().cells
            cells[0].text = str(r["dosen"])
            cells[1].text = str(r["judul"])
            cells[2].text = str(r["tahun"] or "")
            cells[3].text = str(r["sumber"] or "")
        if len(df_pub) > 30:
            doc.add_paragraph(f"... dan {len(df_pub) - 30} baris lainnya "
                               f"(total {len(df_pub)} publikasi tersimpan).").italic = True

    if kode_item == "led_b5_dosen_tendik" and len(df_pub):
        nama_unik = sorted(df_pub["dosen"].unique())
        doc.add_paragraph(f"Jumlah DTPR dengan publikasi terlacak live: {len(nama_unik)} orang "
                           f"(BUKAN hitungan resmi \"jumlah dosen per Des. TS\").")

    if kode_item in ("lkps_2_d", "lkps_4_c_2") and len(df_berita):
        doc.add_paragraph(
            f"{len(df_berita)} berita dcse.fmipa.ugm.ac.id ditandai relevan MEI (perlu verifikasi "
            f"manual satu-satu sebelum dipakai sbg bukti resmi -- lihat data_source_map.json)."
        ).italic = True
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light List Accent 1"
        for i, k in enumerate(["Judul", "Tanggal", "URL"]):
            table.rows[0].cells[i].text = k
        for _, r in df_berita.head(15).iterrows():
            cells = table.add_row().cells
            cells[0].text = str(r["judul"])
            cells[1].text = str(r["tanggal"] or "")
            cells[2].text = str(r["url"])


def _add_cover(doc: Document, judul: str, ringkasan: dict) -> None:
    title = doc.add_heading(judul, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Digenerate otomatis dari data LIVE (Fase 2/3) -- akreditasi/ (UGM Analytics). "
                "BUKAN disalin dari dokumen PDF LED/LKPS lama.").italic = True
    p_tgl = doc.add_paragraph()
    p_tgl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tgl.add_run(f"Digenerate: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph()
    p_stat = doc.add_paragraph()
    p_stat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_stat.add_run(
        f"{ringkasan['tersedia']}/{ringkasan['total']} item tersedia dari sumber live -- "
        f"{ringkasan['akses_data']} tidak tersedia (kendala akses data), "
        f"{ringkasan['perlu_manusia']} perlu disusun tim penyusun (bukan soal akses data)."
    )
    legenda = doc.add_paragraph()
    legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = legenda.add_run("[DATA TIDAK TERSEDIA]")
    r1.font.color.rgb = WARNA_AKSES_DATA
    r1.bold = True
    legenda.add_run(" = sumber ada tapi belum bisa diakses (lihat catatan per item)   |   ")
    r2 = legenda.add_run("[NARASI PERLU DISUSUN TIM PENYUSUN PRODI]")
    r2.font.color.rgb = WARNA_PERLU_MANUSIA
    r2.bold = True
    legenda.add_run(" = di luar cakupan sistem, memang harus ditulis manusia")
    doc.add_page_break()


def _add_item_section(doc: Document, item: dict, map_entry: dict, df_tersedia: pd.DataFrame,
                       df_pub: pd.DataFrame, df_berita: pd.DataFrame) -> None:
    heading_text = item["nama"]
    if item["tabel_lkps"]:
        heading_text = f"Tabel {item['tabel_lkps']} — {item['nama']}"
    doc.add_heading(heading_text, level=2)
    if item["deskripsi_singkat"]:
        desc_p = doc.add_paragraph()
        desc_p.add_run(item["deskripsi_singkat"]).italic = True

    if map_entry["status"] == "tersedia":
        rows = _rows_for_item(item["id"], df_tersedia)
        _add_tersedia(doc, item, rows)
    else:
        _add_placeholder(doc, map_entry)
    _add_sumber_tambahan(doc, map_entry, item["id"], df_pub, df_berita)
    doc.add_paragraph()


def _ringkasan(item_ids: list[str], data_map: dict) -> dict:
    tersedia = sum(1 for i in item_ids if data_map[i]["status"] == "tersedia")
    akses_data = sum(1 for i in item_ids if data_map[i].get("jenis_kendala") == "akses_data")
    perlu_manusia = sum(1 for i in item_ids if data_map[i].get("jenis_kendala") == "perlu_penyusunan_manusia")
    return {"total": len(item_ids), "tersedia": tersedia, "akses_data": akses_data, "perlu_manusia": perlu_manusia}


def generate_led_docx(engine=None) -> bytes:
    engine = engine or db.get_engine()
    data_map = load_map()
    df_tersedia = load_item_tersedia(engine)
    df_pub = load_publikasi_dosen(engine)
    df_berita = load_berita_dcse_relevan(engine)

    grouped = led_items_by_kriteria()
    item_ids = [item["id"] for items in grouped.values() for item in items]
    ringkasan = _ringkasan(item_ids, data_map)

    doc = Document()
    _add_cover(doc, "Laporan Evaluasi Diri (LED) — Data Live", ringkasan)

    for kriteria in URUTAN_KRITERIA:
        items = grouped.get(kriteria, [])
        if not items:
            continue
        doc.add_heading(LABEL_KRITERIA[kriteria], level=1)
        for item in items:
            _add_item_section(doc, item, data_map[item["id"]], df_tersedia, df_pub, df_berita)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_lkps_docx(engine=None) -> bytes:
    engine = engine or db.get_engine()
    data_map = load_map()
    df_tersedia = load_item_tersedia(engine)
    df_pub = load_publikasi_dosen(engine)
    df_berita = load_berita_dcse_relevan(engine)

    grouped = lkps_items_by_bagian()
    item_ids = [item["id"] for items in grouped.values() for item in items]
    ringkasan = _ringkasan(item_ids, data_map)

    doc = Document()
    _add_cover(doc, "Laporan Kinerja Program Studi (LKPS) — Data Live", ringkasan)

    for bagian in URUTAN_BAGIAN_LKPS:
        items = grouped.get(bagian, [])
        if not items:
            continue
        doc.add_heading(LABEL_BAGIAN_LKPS[bagian], level=1)
        for item in items:
            _add_item_section(doc, item, data_map[item["id"]], df_tersedia, df_pub, df_berita)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    out_dir = Path(__file__).resolve().parents[1] / "data" / "generated"
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "Laporan_LED_live_percobaan.docx").write_bytes(generate_led_docx())
    (out_dir / "Laporan_LKPS_live_percobaan.docx").write_bytes(generate_lkps_docx())
    print(f"OK -- {out_dir / 'Laporan_LED_live_percobaan.docx'}")
    print(f"OK -- {out_dir / 'Laporan_LKPS_live_percobaan.docx'}")
