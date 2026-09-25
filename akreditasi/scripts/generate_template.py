"""Generator dokumen template akreditasi (.docx) -- python-docx.

Dua entry point terpisah -- generate_led_docx() dan generate_lkps_docx() --
sesuai pemisahan mode dashboard (LED per Kriteria A-D, LKPS per Bagian 1-6),
lihat docstring registry_kebutuhan_data.py poin 6. Prinsip WAJIB (tidak
berubah): SETIAP item registry di dokumen masing-masing selalu punya
section/tabel, bahkan kalau datanya belum ada -- ditandai placeholder italic
"Data belum tersedia" (bukan section yang dihilangkan), supaya dokumen tetap
berfungsi sebagai template lengkap yang tinggal dilengkapi tim penyusun.

build_led_docx()/build_lkps_docx() menerima DataFrame data_manual yang sudah
dimuat pemanggil (API FastAPI memuatnya dengan SQL portabel MySQL/PostgreSQL),
jadi modul ini tidak wajib menyentuh `db` (koneksi MySQL skrip). `db` hanya
diimpor malas oleh generate_*_docx() untuk pemakaian CLI.

`konteks` (opsional, dikirim API) menggabungkan tiga sumber seperti aturan
generate_laporan_live.py: status tiap item dari data_source_map.json (BUKAN
status_ketersediaan registry), isian tim (data_manual, prioritas utama), lalu
data live pipeline Fase 2 untuk item "tersedia". Item tanpa keduanya diberi
placeholder yang membedakan "[DATA TIDAK TERSEDIA]" (kendala akses data, plus
catatan sumbernya) dari "[NARASI PERLU DISUSUN TIM PENYUSUN PRODI]". Tanpa
`konteks`, perilaku lama (template dari data_manual saja) tidak berubah.
"""

import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from registry_kebutuhan_data import (  # noqa: E402
    KEBUTUHAN_DATA,
    LABEL_BAGIAN_LKPS,
    LABEL_KRITERIA,
    URUTAN_BAGIAN_LKPS,
    URUTAN_KRITERIA,
    led_items_by_kriteria,
    lkps_cuplikan_untuk_kriteria,
    lkps_items_by_bagian,
    ringkasan_status,
)

WARNA_PLACEHOLDER = RGBColor(0x99, 0x33, 0x33)
WARNA_PERLU_MANUSIA = RGBColor(0x1F, 0x4E, 0x79)
WARNA_KETERANGAN = RGBColor(0x55, 0x55, 0x55)


def _load_data_manual(engine, prodi_id: str = "mei") -> pd.DataFrame:
    """`prodi_id` WAJIB di-filter (2026-09-11, lihat scripts/migrasi_prodi_id.py)
    -- tabel data_manual sekarang multi-tenant, tanpa filter ini dokumen yang
    digenerate akan mencampur data semua prodi jadi satu."""
    import db  # malas: lihat docstring modul

    if not db.table_exists(engine, db.t("data_manual")):
        return pd.DataFrame(columns=["item_id", "baris_ke", "kolom", "tahun", "nilai", "link_bukti"])
    return db.read_sql_retry(
        engine, f"SELECT item_id, baris_ke, kolom, tahun, nilai, link_bukti FROM `{db.t('data_manual')}` "
                f"WHERE prodi_id = :prodi_id",
        label="load data_manual (generate_template)", params={"prodi_id": prodi_id},
    )


def _rows_for_item(item_id: str, df: pd.DataFrame) -> list[dict]:
    """Kembalikan list baris (dict kolom->nilai) untuk satu item, dari bentuk
    long (satu baris df = satu sel) dipivot balik ke wide per baris_ke."""
    sub = df[df["item_id"] == item_id]
    if not len(sub):
        return []
    rows: dict[int, dict] = {}
    for _, r in sub.iterrows():
        rows.setdefault(int(r["baris_ke"]), {})[r["kolom"]] = r["nilai"] or ""
        if r.get("link_bukti") and "Link Bukti" not in rows[int(r["baris_ke"])]:
            rows[int(r["baris_ke"])].setdefault("_link_bukti", r["link_bukti"])
    return [rows[k] for k in sorted(rows)]


def _placeholder_paragraph(doc: Document, sumber: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"⚠️ Data belum tersedia — sumber data seharusnya: {sumber}")
    run.italic = True
    run.font.color.rgb = WARNA_PLACEHOLDER


def _add_tabel_item(doc: Document, item: dict, rows: list[dict]) -> None:
    kolom = item["kolom_dibutuhkan"]
    if not rows:
        _placeholder_paragraph(doc, item["sumber_data"])
        return
    table = doc.add_table(rows=1, cols=len(kolom))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, k in enumerate(kolom):
        hdr[i].text = k
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, k in enumerate(kolom):
            cells[i].text = str(row.get(k, ""))


def _add_narasi_item(doc: Document, item: dict, rows: list[dict]) -> None:
    if not rows:
        _placeholder_paragraph(doc, item["sumber_data"])
        return
    row = rows[0]
    for k in item["kolom_dibutuhkan"]:
        nilai = row.get(k, "")
        p = doc.add_paragraph()
        p.add_run(f"{k}: ").bold = True
        if nilai:
            p.add_run(str(nilai))
        else:
            run = p.add_run("⚠️ Data belum tersedia")
            run.italic = True
            run.font.color.rgb = WARNA_PLACEHOLDER


def _add_cover(doc: Document, judul: str, ringkasan: dict) -> None:
    title = doc.add_heading(judul, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Template digenerate otomatis — akreditasi/ (UGM Analytics)").italic = True
    p_tgl = doc.add_paragraph()
    p_tgl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tgl.add_run(f"Digenerate: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph()
    p_stat = doc.add_paragraph()
    p_stat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_stat.add_run(
        f"Kelengkapan data: {ringkasan['lengkap']}/{ringkasan['total']} item "
        f"({100 * ringkasan['lengkap'] / ringkasan['total']:.0f}%) — "
        f"{ringkasan['tersedia_otomatis']} tersedia otomatis, "
        f"{ringkasan['perlu_manual_terisi']}/{ringkasan['perlu_manual_total']} input manual terisi, "
        f"{ringkasan['belum_tersedia']} belum tersedia."
    )
    doc.add_page_break()


def _add_todo(doc: Document, item_ids: list[str], filled_ids: set) -> None:
    doc.add_heading("Ringkasan & To-Do — Item Data Belum Lengkap", level=1)
    doc.add_paragraph(
        "Daftar berikut adalah checklist untuk tim penyusun: seluruh item data yang "
        "masih berstatus belum tersedia atau perlu input manual tapi belum diisi."
    )
    todo_kolom = ["Nama", "Tabel LKPS", "Kriteria", "Sumber Data Seharusnya"]
    todo_table = doc.add_table(rows=1, cols=len(todo_kolom))
    todo_table.style = "Light Grid Accent 1"
    for i, k in enumerate(todo_kolom):
        todo_table.rows[0].cells[i].text = k
        for r in todo_table.rows[0].cells[i].paragraphs[0].runs:
            r.font.bold = True
    for item_id in item_ids:
        item = KEBUTUHAN_DATA[item_id]
        belum_lengkap = (
            item["status_ketersediaan"] == "belum_tersedia"
            or (item["status_ketersediaan"] == "perlu_input_manual" and item_id not in filled_ids)
        )
        if belum_lengkap:
            cells = todo_table.add_row().cells
            cells[0].text = item["nama"]
            cells[1].text = item["tabel_lkps"] or "—"
            cells[2].text = LABEL_KRITERIA.get(item["kriteria_led"], item["kriteria_led"])
            cells[3].text = item["sumber_data"]
    doc.add_page_break()


def _add_item_section(doc: Document, item: dict, df_manual: pd.DataFrame) -> None:
    heading_text = item["nama"]
    if item["tabel_lkps"]:
        heading_text = f"Tabel {item['tabel_lkps']} — {item['nama']}"
    doc.add_heading(heading_text, level=2)
    if item["deskripsi_singkat"]:
        desc_p = doc.add_paragraph()
        desc_p.add_run(item["deskripsi_singkat"]).italic = True
    rows = _rows_for_item(item["id"], df_manual)
    if item["tipe"] == "tabel":
        _add_tabel_item(doc, item, rows)
    else:
        _add_narasi_item(doc, item, rows)
    doc.add_paragraph()


def _add_cuplikan_lkps(doc: Document, kriteria: str, filled_ids: set) -> None:
    """Sub-tabel ringkas (nama, no. tabel, status saja -- BUKAN isi penuh) di
    dokumen LED, dipakai untuk Kriteria A/B/C1-C6 -- isi lengkap ada di
    dokumen LKPS, jangan duplikasi data (lihat registry poin 6)."""
    cuplikan = lkps_cuplikan_untuk_kriteria(kriteria)
    if not cuplikan:
        return
    doc.add_heading("Tabel LKPS Terkait (ringkasan)", level=3)
    doc.add_paragraph(
        "Isi lengkap tabel-tabel berikut ada di Dokumen LKPS -- dicantumkan di sini "
        "sebagai ringkasan bukti evaluasi PPEPP."
    ).italic = True
    kolom = ["Tabel LKPS", "Nama", "Status"]
    table = doc.add_table(rows=1, cols=len(kolom))
    table.style = "Light Grid Accent 1"
    for i, k in enumerate(kolom):
        table.rows[0].cells[i].text = k
        for r in table.rows[0].cells[i].paragraphs[0].runs:
            r.font.bold = True
    for item in cuplikan:
        cells = table.add_row().cells
        cells[0].text = item["tabel_lkps"]
        cells[1].text = item["nama"]
        cells[2].text = "Terisi" if item["id"] in filled_ids else "Belum terisi"
    doc.add_paragraph()


def generate_led_docx(engine=None, prodi_id: str = "mei") -> bytes:
    """Muat data_manual lewat `db` (MySQL skrip) lalu bangun dokumen LED -- untuk CLI."""
    import db  # malas: lihat docstring modul

    return build_led_docx(_load_data_manual(engine or db.get_engine(), prodi_id))


def build_led_docx(df_manual: pd.DataFrame, konteks: dict | None = None) -> bytes:
    """Dokumen "Laporan Evaluasi Diri (LED)" -- isi per Kriteria (Umum/A/B/C1-C6/D),
    lihat led_items_by_kriteria(). Kriteria A/B/C1-C6 ditutup sub-tabel ringkas
    tabel LKPS terkait (cuplikan, bukan isi penuh). `df_manual` = baris
    akreditasi_data_manual SATU prodi (kolom item_id, baris_ke, kolom, nilai, link_bukti)."""
    if konteks is not None:
        return _build_gabungan("Laporan Evaluasi Diri (LED)", led_items_by_kriteria(), URUTAN_KRITERIA,
                               LABEL_KRITERIA, df_manual, konteks, cuplikan=True)
    filled_ids = set(df_manual["item_id"].unique()) if len(df_manual) else set()
    grouped = led_items_by_kriteria()
    item_ids = [item["id"] for items in grouped.values() for item in items]
    ringkasan = ringkasan_status(item_ids=item_ids, terisi_ids=filled_ids)

    doc = Document()
    _add_cover(doc, "Laporan Evaluasi Diri (LED)", ringkasan)
    _add_todo(doc, item_ids, filled_ids)

    for kriteria in URUTAN_KRITERIA:
        items = grouped.get(kriteria, [])
        if not items:
            continue
        doc.add_heading(LABEL_KRITERIA[kriteria], level=1)
        for item in items:
            _add_item_section(doc, item, df_manual)
        if kriteria not in ("Umum", "D"):
            _add_cuplikan_lkps(doc, kriteria, filled_ids)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_lkps_docx(engine=None, prodi_id: str = "mei") -> bytes:
    """Muat data_manual lewat `db` (MySQL skrip) lalu bangun dokumen LKPS -- untuk CLI."""
    import db  # malas: lihat docstring modul

    return build_lkps_docx(_load_data_manual(engine or db.get_engine(), prodi_id))


def build_lkps_docx(df_manual: pd.DataFrame, konteks: dict | None = None) -> bytes:
    """Dokumen "Laporan Kinerja Program Studi (LKPS)" -- isi per Bagian 1-6,
    lihat lkps_items_by_bagian(). Rendering tabel sama persis seperti
    sebelumnya, cuma dikelompok ulang per Bagian bukan per Kriteria."""
    if konteks is not None:
        return _build_gabungan("Laporan Kinerja Program Studi (LKPS)", lkps_items_by_bagian(), URUTAN_BAGIAN_LKPS,
                               LABEL_BAGIAN_LKPS, df_manual, konteks, cuplikan=False)
    filled_ids = set(df_manual["item_id"].unique()) if len(df_manual) else set()
    grouped = lkps_items_by_bagian()
    item_ids = [item["id"] for items in grouped.values() for item in items]
    ringkasan = ringkasan_status(item_ids=item_ids, terisi_ids=filled_ids)

    doc = Document()
    _add_cover(doc, "Laporan Kinerja Program Studi (LKPS)", ringkasan)
    _add_todo(doc, item_ids, filled_ids)

    for bagian in URUTAN_BAGIAN_LKPS:
        items = grouped.get(bagian, [])
        if not items:
            continue
        doc.add_heading(LABEL_BAGIAN_LKPS[bagian], level=1)
        for item in items:
            _add_item_section(doc, item, df_manual)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- mode gabungan (konteks)
def _kategori(peta: dict, item_id: str) -> str:
    entri = peta.get(item_id) or {}
    if entri.get("status") == "tersedia":
        return "tersedia"
    return "penyusunan" if entri.get("jenis_kendala") == "perlu_penyusunan_manusia" else "akses_data"


def _keterangan(doc: Document, teks: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(teks)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = WARNA_KETERANGAN


def _placeholder_peta(doc: Document, peta: dict, item_id: str) -> None:
    if _kategori(peta, item_id) == "penyusunan":
        run = doc.add_paragraph().add_run("[NARASI PERLU DISUSUN TIM PENYUSUN PRODI]")
        run.bold = True
        run.font.color.rgb = WARNA_PERLU_MANUSIA
        return
    run = doc.add_paragraph().add_run("[DATA TIDAK TERSEDIA]")
    run.bold = True
    run.font.color.rgb = WARNA_PLACEHOLDER
    catatan = (peta.get(item_id) or {}).get("catatan_jika_tidak_tersedia")
    if catatan:
        p = doc.add_paragraph()
        p.add_run("Catatan: ").bold = True
        p.add_run(catatan)


def _add_live(doc: Document, item: dict, live_item: dict) -> None:
    """Data live item "tersedia": semua kolom resmi tetap tampil; sel yang tidak berhasil ditarik
    ditandai [DATA TIDAK TERSEDIA] (bukan dihilangkan), kolom tambahan dari pipeline ikut di akhir."""
    rows = live_item["rows"]
    kolom = list(item["kolom_dibutuhkan"])
    for row in rows:
        kolom += [k for k in row if k not in kolom]
    if len(rows) == 1 and len(kolom) <= 8:
        for k in kolom:
            p = doc.add_paragraph()
            p.add_run(f"{k}: ").bold = True
            nilai = rows[0].get(k, "")
            run = p.add_run(str(nilai) if nilai else "[DATA TIDAK TERSEDIA]")
            if not nilai:
                run.font.color.rgb = WARNA_PLACEHOLDER
    else:
        table = doc.add_table(rows=1, cols=len(kolom))
        table.style = "Light Grid Accent 1"
        for i, k in enumerate(kolom):
            table.rows[0].cells[i].text = k
            for r in table.rows[0].cells[i].paragraphs[0].runs:
                r.font.bold = True
        for row in rows:
            cells = table.add_row().cells
            for i, k in enumerate(kolom):
                cells[i].text = str(row.get(k, "")) or "[DATA TIDAK TERSEDIA]"
    sumber = "; ".join(live_item.get("sumber") or [])
    tanggal = str(live_item.get("fetched_at") or "")[:10]
    _keterangan(doc, f"Sumber live: {sumber or '-'}" + (f" (diambil {tanggal})" if tanggal else "")
                + ". Periksa ulang sebelum dipakai sebagai data resmi.")


def _add_pendukung(doc: Document, peta: dict, item_id: str, live: dict) -> None:
    teks = (peta.get(item_id) or {}).get("sumber_tambahan")
    publikasi, berita = live.get("publikasi") or [], live.get("berita") or []
    if not teks:
        return
    tabel: tuple[list[str], list[list[str]], int] | None = None
    if item_id in ("lkps_3_c_2", "lkps_3_a_2") and publikasi:
        tabel = (["Dosen", "Judul", "Tahun", "Sumber"],
                 [[str(r["dosen"]), str(r["judul"]), str(r["tahun"] or ""), str(r["sumber"] or "")]
                  for r in publikasi[:30]], len(publikasi))
    elif item_id == "led_b5_dosen_tendik" and publikasi:
        dosen = sorted({r["dosen"] for r in publikasi})
        tabel = (["Dosen dengan publikasi terlacak (SINTA)"], [[d] for d in dosen], len(dosen))
    elif item_id in ("lkps_2_d", "lkps_4_c_2") and berita:
        tabel = (["Judul", "Tanggal", "URL"],
                 [[str(r["judul"]), str(r["tanggal"] or "")[:10], str(r["url"])] for r in berita[:15]], len(berita))
    if tabel is None:
        return
    doc.add_heading("Data pendukung (live, bukan pengganti isian resmi di atas)", level=3)
    _keterangan(doc, teks)
    kolom, rows, total = tabel
    table = doc.add_table(rows=1, cols=len(kolom))
    table.style = "Light List Accent 1"
    for i, k in enumerate(kolom):
        table.rows[0].cells[i].text = k
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    if total > len(rows):
        _keterangan(doc, f"Ditampilkan {len(rows)} dari {total} baris.")


def _add_item_gabungan(doc: Document, item: dict, df_manual: pd.DataFrame, konteks: dict) -> None:
    heading_text = f"Tabel {item['tabel_lkps']} — {item['nama']}" if item["tabel_lkps"] else item["nama"]
    doc.add_heading(heading_text, level=2)
    if item["deskripsi_singkat"]:
        doc.add_paragraph().add_run(item["deskripsi_singkat"]).italic = True
    peta, live = konteks["peta"], konteks["live"]
    rows = _rows_for_item(item["id"], df_manual)
    live_item = live["items"].get(item["id"])
    if rows:
        (_add_tabel_item if item["tipe"] == "tabel" else _add_narasi_item)(doc, item, rows)
        _keterangan(doc, "Sumber: isian tim penyusun.")
    elif live_item:
        _add_live(doc, item, live_item)
    else:
        _placeholder_peta(doc, peta, item["id"])
    _add_pendukung(doc, peta, item["id"], live)
    doc.add_paragraph()


def _ringkasan_gabungan(item_ids: list[str], df_manual: pd.DataFrame, konteks: dict) -> dict:
    peta, live_ids = konteks["peta"], set(konteks["live"]["items"])
    terisi = (set(df_manual["item_id"].unique()) if len(df_manual) else set()) & set(item_ids)
    lengkap = {i for i in item_ids if i in terisi or i in live_ids}
    per = {k: [i for i in item_ids if _kategori(peta, i) == k] for k in ("tersedia", "akses_data", "penyusunan")}
    return {"total": len(item_ids), "lengkap": lengkap, "terisi": terisi, "live": live_ids & set(item_ids),
            "per": {k: (sum(1 for i in v if i in lengkap), len(v)) for k, v in per.items()}}


def _add_cover_gabungan(doc: Document, judul: str, r: dict, konteks: dict) -> None:
    title = doc.add_heading(judul, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for teks in (konteks.get("prodi"), konteks.get("fakultas")):
        if teks:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(str(teks)).bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Digenerate: {datetime.now().strftime('%Y-%m-%d %H:%M')} — UGM Analytics").italic = True
    doc.add_paragraph()
    n = len(r["lengkap"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t, a, m = r["per"]["tersedia"], r["per"]["akses_data"], r["per"]["penyusunan"]
    p.add_run(
        f"Kelengkapan data: {n}/{r['total']} item ({100 * n / r['total']:.0f}%) — "
        f"{len(r['terisi'])} diisi tim, {len(r['live'] - r['terisi'])} dari data live. "
        f"Per status sumber: tersedia live {t[0]}/{t[1]}, perlu akses data {a[0]}/{a[1]}, "
        f"perlu disusun tim {m[0]}/{m[1]}."
    )
    legenda = doc.add_paragraph()
    legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = legenda.add_run("[DATA TIDAK TERSEDIA]")
    r1.bold, r1.font.color.rgb = True, WARNA_PLACEHOLDER
    legenda.add_run(" = sumber ada tetapi belum bisa diakses sistem   |   ")
    r2 = legenda.add_run("[NARASI PERLU DISUSUN TIM PENYUSUN PRODI]")
    r2.bold, r2.font.color.rgb = True, WARNA_PERLU_MANUSIA
    legenda.add_run(" = memang harus ditulis tim penyusun")
    doc.add_page_break()


def _add_todo_gabungan(doc: Document, item_ids: list[str], r: dict, konteks: dict) -> None:
    doc.add_heading("Ringkasan & To-Do — Item Data Belum Lengkap", level=1)
    doc.add_paragraph("Checklist tim penyusun: item yang belum punya isian tim maupun data live.")
    label = {"tersedia": "Tersedia live (pipeline belum berhasil)", "akses_data": "Perlu akses data",
             "penyusunan": "Perlu disusun tim"}
    kolom = ["Nama", "Tabel LKPS", "Kriteria", "Status sumber", "Catatan"]
    table = doc.add_table(rows=1, cols=len(kolom))
    table.style = "Light Grid Accent 1"
    for i, k in enumerate(kolom):
        table.rows[0].cells[i].text = k
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.font.bold = True
    peta = konteks["peta"]
    for item_id in item_ids:
        if item_id in r["lengkap"]:
            continue
        item = KEBUTUHAN_DATA[item_id]
        cells = table.add_row().cells
        cells[0].text = item["nama"]
        cells[1].text = item["tabel_lkps"] or "—"
        cells[2].text = LABEL_KRITERIA.get(item["kriteria_led"], item["kriteria_led"])
        cells[3].text = label[_kategori(peta, item_id)]
        cells[4].text = (peta.get(item_id) or {}).get("catatan_jika_tidak_tersedia") or item["sumber_data"]
    doc.add_page_break()


def _build_gabungan(judul: str, grouped: dict, urutan: list[str], label: dict, df_manual: pd.DataFrame,
                    konteks: dict, cuplikan: bool) -> bytes:
    item_ids = [item["id"] for items in grouped.values() for item in items]
    r = _ringkasan_gabungan(item_ids, df_manual, konteks)
    doc = Document()
    _add_cover_gabungan(doc, judul, r, konteks)
    _add_todo_gabungan(doc, item_ids, r, konteks)
    lengkap_semua = (set(df_manual["item_id"].unique()) if len(df_manual) else set()) | set(konteks["live"]["items"])
    for key in urutan:
        items = grouped.get(key, [])
        if not items:
            continue
        doc.add_heading(label[key], level=1)
        for item in items:
            _add_item_gabungan(doc, item, df_manual, konteks)
        if cuplikan and key not in ("Umum", "D"):
            _add_cuplikan_lkps(doc, key, lengkap_semua)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    out_dir = Path(__file__).resolve().parents[1]
    (out_dir / "template_LED.docx").write_bytes(generate_led_docx())
    (out_dir / "template_LKPS.docx").write_bytes(generate_lkps_docx())
    print(f"OK -- {out_dir / 'template_LED.docx'}")
    print(f"OK -- {out_dir / 'template_LKPS.docx'}")
